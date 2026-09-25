"""
Excel export of the Material Take-Off (quantities only - no costs).

Design goals (v0.5.1):
* A non-technical reader can open the Summary and see WHAT to buy and HOW MUCH:
  a shopping list grouped by trade (collapsible), in purchase units
  (bags, tons, coils, pipe lengths, bricks, cft ...), with when it is needed
  and whether the quantity is firm or still based on defaults.
* Only materials inside the selected scope are exported.
* Interactive: contents with hyperlinks, "Back to Summary" links on every
  sheet, Excel Tables (filter/sort buttons, banded rows), collapsible
  trade/stage groups, colour-coded status, click a Mat_ID to jump to its
  calculation breakdown, editable wastage flowing into all quantities.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from detailed_mto.engine import (ST_CALC, ST_CALC_ASSUMED, ST_NEEDS_INPUT, ST_NOT_REQ, ST_OPTION, ST_PROVISIONAL,
                                 ST_REFERENCE, DetailedResult, MaterialLine, purchase_rule)

FONT = "Arial"
NAVY, TEAL = "1F3864", "0D9488"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
CAT_FILL = PatternFill("solid", fgColor="D9E1F2")
CARD_FILL = PatternFill("solid", fgColor="E8F3F1")
INPUT_FONT_COLOR = "0000FF"
LINK_COLOR = "0563C1"
STATUS_COLORS = {
    ST_CALC: "E2EFDA", ST_CALC_ASSUMED: "FFF2CC", ST_PROVISIONAL: "FCE4D6", ST_NEEDS_INPUT: "F8CBAD",
    ST_OPTION: "EDEDED", ST_NOT_REQ: "F2F2F2", ST_REFERENCE: "DDEBF7",
}
thin = Side(style="thin", color="D0D7E2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
NOWRAP = Alignment(wrap_text=False, vertical="top")  # long trace text: keeps rows compact, full text in the cell

STAGE_NAMES = {
    "S0": "Before construction", "S1": "Foundations & substructure", "S2": "Plumbing/electric rough-in",
    "S3": "Brick walls", "S4": "RCC columns, beams & slabs", "S5": "Pipes & wiring inside walls",
    "S6": "Plaster, tiles, ceilings & paint", "S7": "Roof & outside walls", "S8": "Outside works & connections",
    "S9": "Fixtures, fittings & handover",
}

BENCHMARKS = [  # (label, material ids, unit, low, high)
    ("Cement per sft covered area", ["CON-001"], "bags/sft", 0.30, 0.55),
    ("Steel per sft covered area", ["RBR-001", "RBR-002", "RBR-003", "RBR-004"], "kg/sft", 1.5, 2.6),
    ("Bricks per sft covered area", ["MAS-001", "MAS-002"], "Nos/sft", 20, 35),
    ("Sand per sft covered area", ["CON-004", "CON-005"], "cft/sft", 1.2, 2.2),
    ("Crush per sft covered area", ["CON-006", "CON-007", "CON-008"], "cft/sft", 0.7, 1.4),
]

KEY_CARDS = [  # (label, ids, unit, divisor, roundup digits)
    ("Cement", ["CON-001"], "bags", 1, 0),
    ("Steel (all sizes)", ["RBR-001", "RBR-002", "RBR-003", "RBR-004"], "ton", 1000, 2),
    ("Bricks", ["MAS-001", "MAS-002"], "Nos", 1, -2),
    ("Sand", ["CON-004", "CON-005", "EW-003", "EXT-003", "PDR-020"], "cft", 1, -1),
    ("Crush / aggregate", ["CON-006", "CON-007", "CON-008"], "cft", 1, -1),
    ("Floor & wall tiles", ["FLR-001", "FLR-002", "FLR-003"], "sft", 1, 0),
    ("Paint (interior + exterior)", ["PNT-003", "PNT-004", "PNT-005"], "gal", 1, 0),
    ("House wiring (all sizes)", ["ELE-006", "ELE-007", "ELE-008", "ELE-009", "ELE-011"], "coils (90 m)", 90, 0),
]


def _app_version() -> str:
    try:
        import config
        return config.APP_VERSION
    except Exception:  # noqa: BLE001
        return "?"


def _kb_id() -> str:
    try:
        from knowledge import load_knowledge_base
        kb = load_knowledge_base()
        return f"v{kb.version} #{kb.file_hash}"
    except Exception:  # noqa: BLE001
        return "?"


def benchmarks_for(project):
    """Load-bearing houses (strip foundations) carry less steel than RCC frames."""
    out = []
    frame = project.v("FDN_STRIP", 1.0) < 0.5
    for label, ids, unit, lo, hi in BENCHMARKS:
        if unit == "kg/sft" and frame:
            label, lo, hi = label + " (RCC frame)", 2.5, 3.6
        elif unit == "kg/sft":
            label = label + " (load-bearing)"
        out.append((label, ids, unit, lo, hi))
    return out


# ---------------------------------------------------------------------------
def _font(size=9, bold=False, color=None, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def _txt(v) -> str:
    s = "" if v is None else str(v)
    return ("'" + s) if s.startswith("=") else s


def _cell(ws, r, c, v, bold=False, fill=None, fmt=None, color=None, align=WRAP, border=True, size=9):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = _font(size=size, bold=bold, color=color)
    cell.alignment = align
    if border:
        cell.border = BORDER
    if fill is not None:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def _link(ws, r, c, text, target, size=9, bold=False, border=False):
    cell = ws.cell(row=r, column=c, value=text)
    cell.hyperlink = f"#{target}"
    cell.font = _font(size=size, bold=bold, color=LINK_COLOR, underline="single")
    cell.alignment = Alignment(vertical="top")
    if border:
        cell.border = BORDER
    return cell


def _titled_sheet(wb, name, title, subtitle, headers, widths):
    """Row 1 title, row 2 back-link + subtitle, row 3 table header, data from row 4."""
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value=title).font = _font(size=13, bold=True, color=NAVY)
    _link(ws, 2, 1, "◄ Summary", "'Summary'!A1")
    ws.cell(row=2, column=2, value=subtitle).font = _font(italic=True, color="595959")
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1] if i - 1 < len(widths) else 14
    ws.row_dimensions[3].height = 30
    ws.freeze_panes = "B4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "3:3"
    return ws


def _add_table(ws, name, ncols, last_row, style="TableStyleLight9"):
    if last_row < 4:
        return
    tbl = Table(displayName=name, ref=f"A3:{get_column_letter(ncols)}{last_row}")
    tbl.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True, showColumnStripes=False)
    ws.add_table(tbl)


def _status_formatting(ws, col, first, last):
    if last < first:
        return
    for st, color in STATUS_COLORS.items():
        ws.conditional_formatting.add(
            f"{col}{first}:{col}{last}",
            FormulaRule(formula=[f'${col}{first}="{st}"'], fill=PatternFill("solid", fgColor=color), stopIfTrue=True))


def _stage_code(m: MaterialLine) -> str:
    return (m.material.stage or "S?").split(",")[0].split("-")[0].strip()


def _friendly_conf(m: MaterialLine) -> str:
    if m.material.unit.upper() == "LS":
        return "Lump-sum item"
    if m.status == ST_PROVISIONAL:
        return "Allowance - confirm with owner"
    if m.confidence in ("High", "Medium", "User"):
        return "✔ From drawings / your inputs"
    return "⚠ Check - uses default values"


def _buy_formula(m: MaterialLine, src: str) -> Tuple[str, str]:
    """Excel expression (without '=') for the buy quantity + buy unit."""
    div, digits, unit = purchase_rule(m.material)
    if div is None:
        return f"IF({src}>0,1,0)", unit
    inner = src if div == 1 else f"{src}/{div}"
    return f"ROUNDUP({inner},{digits})", unit


def _buy_fmt(unit: str) -> str:
    return "#,##0.00" if unit == "ton" else "#,##0"


# ---------------------------------------------------------------------------
def build_detailed_mto_workbook(result: DetailedResult) -> bytes:
    p = result.project
    mats = result.scoped_materials()  # out-of-scope materials are never exported
    wb = Workbook()
    wb.remove(wb.active)
    summary = wb.create_sheet("Summary")

    # ============================== Material_Schedule =========================
    MH = ["Mat_ID", "Category", "Subcategory", "Material Description", "Specification/Grade", "Unit", "Net Qty",
          "Wastage %", "Qty incl. Wastage", "Buy Qty", "Buy Unit", "Status", "Confidence", "Stage", "Tier/Option",
          "Calculation (traceable)", "Driven by Work Items", "Required Inputs", "Notes", "If-selected / Reference Qty"]
    ws = _titled_sheet(wb, "Material_Schedule", "Material Schedule - every material in scope",
                       f"Scope: {result.scope}. Blue wastage % is editable - quantities and the Summary update. "
                       "Click a Mat_ID to see how it was calculated. Use the filter buttons to sort/filter.",
                       MH, [9, 20, 15, 38, 26, 6, 11, 8, 12, 10, 14, 17, 10, 7, 10, 60, 18, 24, 34, 12])
    sched_row: Dict[str, int] = {}
    contribs = sorted(result.scoped_contributions(), key=lambda c: (c.mat_id, c.wi_id))
    breakdown_first: Dict[str, int] = {}
    for i, c in enumerate(contribs):
        breakdown_first.setdefault(c.mat_id, i + 4)
    for i, m in enumerate(mats):
        r = i + 4
        mat = m.material
        sched_row[mat.mat_id] = r
        if mat.mat_id in breakdown_first:
            _link(ws, r, 1, mat.mat_id, f"'Material_Breakdown'!A{breakdown_first[mat.mat_id]}", border=True)
        else:
            _cell(ws, r, 1, mat.mat_id)
        for j, v in enumerate([mat.category, mat.subcategory, mat.description, mat.specification, mat.unit], 2):
            _cell(ws, r, j, _txt(v))
        _cell(ws, r, 7, round(m.net_qty, 3), fmt="#,##0.00")
        _cell(ws, r, 8, m.wastage_pct, fmt="0.0%", color=INPUT_FONT_COLOR)
        _cell(ws, r, 9, f"=G{r}*(1+H{r})", fmt="#,##0.00")
        f, unit = _buy_formula(m, f"I{r}")
        _cell(ws, r, 10, f"=IF(I{r}>0,{f},0)", fmt=_buy_fmt(unit), bold=True)
        _cell(ws, r, 11, unit)
        _cell(ws, r, 12, m.status)
        _cell(ws, r, 13, m.confidence)
        _cell(ws, r, 14, _stage_code(m))
        _cell(ws, r, 15, mat.tier)
        _cell(ws, r, 16, _txt(m.calculation), align=NOWRAP)
        _cell(ws, r, 17, ", ".join(m.driven_by), align=NOWRAP)
        _cell(ws, r, 18, _txt(mat.required_inputs), align=NOWRAP)
        _cell(ws, r, 19, _txt(mat.notes), align=NOWRAP)
        _cell(ws, r, 20, round(m.if_selected_qty, 3) if m.if_selected_qty else None, fmt="#,##0.00")
    last = max(len(mats) + 3, 4)
    _add_table(ws, "MaterialSchedule", len(MH), last)
    _status_formatting(ws, "L", 4, last)
    ws.freeze_panes = "E4"
    SCH_ID = f"Material_Schedule!$A$4:$A${last}"
    SCH_QTY = f"Material_Schedule!$I$4:$I${last}"

    # ============================== Procurement_by_Stage ======================
    ws2 = _titled_sheet(wb, "Procurement_by_Stage", "What to buy, stage by stage",
                        "Grouped by construction stage - use the +/- buttons at the left edge to open or close a stage.",
                        ["Stage / Mat_ID", "Material", "Specification", "Buy Qty", "Buy Unit", "Qty incl. Wastage", "Unit", "Status"],
                        [14, 40, 30, 11, 16, 13, 8, 20])
    ws2.freeze_panes = "A4"
    by_stage: Dict[str, List[MaterialLine]] = {}
    for m in result.purchase_list():
        if m.material.mat_id in sched_row:
            by_stage.setdefault(_stage_code(m), []).append(m)
    r = 4
    for stg in sorted(by_stage):
        _cell(ws2, r, 1, stg, bold=True, fill=CAT_FILL)
        _cell(ws2, r, 2, f"{STAGE_NAMES.get(stg, '')}  ({len(by_stage[stg])} items)", bold=True, fill=CAT_FILL)
        for c in range(3, 9):
            _cell(ws2, r, c, None, fill=CAT_FILL)
        r += 1
        for m in sorted(by_stage[stg], key=lambda x: (x.material.category, x.material.mat_id)):
            sr = sched_row[m.material.mat_id]
            _, unit = _buy_formula(m, "x")
            _link(ws2, r, 1, m.material.mat_id, f"'Material_Schedule'!A{sr}", border=True)
            _cell(ws2, r, 2, m.material.description)
            _cell(ws2, r, 3, _txt(m.material.specification))
            _cell(ws2, r, 4, f"=Material_Schedule!J{sr}", fmt=_buy_fmt(unit), bold=True)
            _cell(ws2, r, 5, unit)
            _cell(ws2, r, 6, f"=Material_Schedule!I{sr}", fmt="#,##0.00")
            _cell(ws2, r, 7, m.material.unit)
            _cell(ws2, r, 8, m.status)
            ws2.row_dimensions[r].outlineLevel = 1
            r += 1
    _status_formatting(ws2, "H", 4, r - 1)
    ws2.sheet_properties.outlinePr.summaryBelow = False

    # ============================== BOQ_Work_Items ============================
    wis = result.scoped_work_items()
    ws3 = _titled_sheet(wb, "BOQ_Work_Items", "Measured work items (quantities only)",
                        "Each work item is measured from the drawing inputs; its materials come from the recipes.",
                        ["WI_ID", "Scope", "Division", "Work Item", "Unit", "Quantity", "Status", "Confidence",
                         "Calculation (from drawing inputs)", "Measurement Rule"],
                        [10, 11, 14, 40, 6, 11, 18, 10, 70, 44])
    for i, w in enumerate(wis):
        r = i + 4
        for j, v in enumerate([w.wi_id, w.scope, w.division, w.description, w.unit], 1):
            _cell(ws3, r, j, _txt(v))
        _cell(ws3, r, 6, round(w.qty, 3), fmt="#,##0.00")
        _cell(ws3, r, 7, w.status)
        _cell(ws3, r, 8, w.confidence)
        _cell(ws3, r, 9, _txt(w.calculation), align=NOWRAP)
        _cell(ws3, r, 10, _txt(w.measurement_rule), align=NOWRAP)
    _add_table(ws3, "WorkItems", 10, len(wis) + 3)
    _status_formatting(ws3, "G", 4, len(wis) + 3)

    # ============================== Material_Breakdown ========================
    ws4 = _titled_sheet(wb, "Material_Breakdown", "How each material quantity was calculated",
                        "Material qty = work-item qty x coefficient (from the Master Material Database). Click Mat_ID to go back.",
                        ["Mat_ID", "Material", "WI_ID", "Work Item", "WI Qty", "WI Unit", "Coefficient", "Coeff_ID", "Mix",
                         "Net Material Qty", "Material Unit"],
                        [10, 36, 10, 36, 11, 7, 10, 20, 12, 13, 8])
    mat_by_id = {m.material.mat_id: m for m in mats}
    for i, c in enumerate(contribs):
        r = i + 4
        m = mat_by_id.get(c.mat_id)
        if c.mat_id in sched_row:
            _link(ws4, r, 1, c.mat_id, f"'Material_Schedule'!A{sched_row[c.mat_id]}", border=True)
        else:
            _cell(ws4, r, 1, c.mat_id)
        _cell(ws4, r, 2, m.material.description if m else "")
        _cell(ws4, r, 3, c.wi_id)
        _cell(ws4, r, 4, c.wi_desc)
        _cell(ws4, r, 5, round(c.wi_qty, 3), fmt="#,##0.00")
        _cell(ws4, r, 6, c.wi_unit)
        _cell(ws4, r, 7, round(c.coefficient, 6), fmt="0.0000", color=INPUT_FONT_COLOR)
        _cell(ws4, r, 8, c.coeff_id)
        _cell(ws4, r, 9, c.mix_ref)
        _cell(ws4, r, 10, f"=E{r}*G{r}", fmt="#,##0.00")
        _cell(ws4, r, 11, m.material.unit if m else "")
    _add_table(ws4, "Breakdown", 11, len(contribs) + 3)

    # ============================== Project_Inputs ============================
    params = sorted(p.params.values(), key=lambda x: (x.group, x.key))
    ws5 = _titled_sheet(wb, "Project_Inputs", "Every input used, with where it came from",
                        "Source shows the drawing sheet, derivation rule or default. Confidence 'Assumed' = please verify.",
                        ["Key", "Group", "Parameter", "Value", "Unit", "Source", "Confidence", "Note"],
                        [16, 12, 44, 11, 7, 46, 10, 50])
    for i, prm in enumerate(params):
        r = i + 4
        for j, v in enumerate([prm.key, prm.group, prm.label], 1):
            _cell(ws5, r, j, v)
        _cell(ws5, r, 4, round(prm.value, 3), fmt="#,##0.00", color=INPUT_FONT_COLOR)
        _cell(ws5, r, 5, prm.unit)
        _cell(ws5, r, 6, _txt(prm.source))
        _cell(ws5, r, 7, prm.confidence)
        _cell(ws5, r, 8, _txt(prm.note))
    _add_table(ws5, "Inputs", 8, len(params) + 3)
    base = len(params) + 5
    _cell(ws5, base, 1, "Floors", bold=True, fill=CAT_FILL)
    for j, h in enumerate(["Floor", "Covered (sft)", "Ext. perimeter (rft)", "9in walls (rft)", "4.5in walls (rft)",
                           "Storey height (ft)", "Source", "Confidence"], 1):
        _cell(ws5, base + 1, j, h, bold=True, fill=CAT_FILL)
    for i, fl in enumerate(p.floors):
        rr = base + 2 + i
        for j, v in enumerate([fl.name, fl.covered_sft, fl.ext_perimeter_ft, fl.wall9_len_ft, fl.wall45_len_ft,
                               fl.storey_height_ft, fl.source, fl.confidence], 1):
            _cell(ws5, rr, j, round(v, 2) if isinstance(v, float) else v, fmt="#,##0.0" if isinstance(v, float) else None)
    opt_r = base + 3 + len(p.floors)
    _cell(ws5, opt_r, 1, "Take-off options", bold=True, fill=CAT_FILL)
    for i, (k, v) in enumerate(vars(p.options).items()):
        _cell(ws5, opt_r + 1 + i, 1, k)
        _cell(ws5, opt_r + 1 + i, 3, str(v))

    # ============================== Rooms =====================================
    ws6 = _titled_sheet(wb, "Rooms", "Rooms used for finishes",
                        "Room sizes as used in the take-off (read from the drawings or edited in the app).",
                        ["Floor", "Room (label)", "Room Type", "Length (ft)", "Width (ft)", "Area (sft)", "Perimeter (rft)",
                         "Wet area?", "Source", "Confidence"], [10, 22, 22, 9, 9, 10, 10, 8, 30, 10])
    for i, rm in enumerate(p.rooms):
        r = i + 4
        _cell(ws6, r, 1, rm.floor)
        _cell(ws6, r, 2, rm.name)
        _cell(ws6, r, 3, rm.room_type)
        _cell(ws6, r, 4, round(rm.length_ft, 3), fmt="0.00")
        _cell(ws6, r, 5, round(rm.width_ft, 3), fmt="0.00")
        _cell(ws6, r, 6, f"=D{r}*E{r}", fmt="#,##0.0")
        _cell(ws6, r, 7, f"=2*(D{r}+E{r})", fmt="#,##0.0")
        _cell(ws6, r, 8, "Yes" if rm.is_wet else "No")
        _cell(ws6, r, 9, rm.source)
        _cell(ws6, r, 10, rm.confidence)
    lr = len(p.rooms) + 3
    _add_table(ws6, "Rooms", 10, lr)
    _cell(ws6, lr + 2, 2, "TOTAL floor area", bold=True)
    _cell(ws6, lr + 2, 6, f"=SUM(F4:F{lr})", bold=True, fmt="#,##0")

    # ============================== Openings ==================================
    ws7 = _titled_sheet(wb, "Openings", "Doors, windows & ventilators",
                        "Opening sizes as used in the take-off.",
                        ["Kind", "Name", "Width (ft)", "Height (ft)", "Qty", "Leaves", "Chogath width (in)", "External?",
                         "Area each (sft)", "Total area (sft)", "Source", "Confidence"],
                        [10, 26, 9, 9, 6, 7, 9, 8, 10, 11, 26, 10])
    for i, o in enumerate(p.openings):
        r = i + 4
        _cell(ws7, r, 1, o.kind)
        _cell(ws7, r, 2, o.name)
        _cell(ws7, r, 3, o.width_ft, fmt="0.00")
        _cell(ws7, r, 4, o.height_ft, fmt="0.00")
        _cell(ws7, r, 5, o.qty)
        _cell(ws7, r, 6, o.leaves)
        _cell(ws7, r, 7, o.chogath_in or None)
        _cell(ws7, r, 8, "Yes" if o.external else "No")
        _cell(ws7, r, 9, f"=C{r}*D{r}", fmt="0.0")
        _cell(ws7, r, 10, f"=I{r}*E{r}", fmt="#,##0.0")
        _cell(ws7, r, 11, o.source)
        _cell(ws7, r, 12, o.confidence)
    _add_table(ws7, "Openings", 12, len(p.openings) + 3)

    # ============================== Assumptions_Gaps ==========================
    items = [("Drawing conflict", c) for c in p.conflicts] + [("Assumption", a) for a in p.assumptions]
    for m in mats:
        if m.status == ST_NEEDS_INPUT:
            items.append(("Needs input", f"{m.material.mat_id} {m.material.description}: provide {m.material.required_inputs or 'quantity'}"))
    for prm in p.params.values():
        if prm.confidence in ("Assumed", "Low"):
            items.append(("Default value used", f"{prm.label} = {prm.value:,.2f} {prm.unit} ({prm.source}) - please verify"))
    ws8 = _titled_sheet(wb, "Assumptions_Gaps", "Things to check before buying",
                        "Drawing conflicts, assumptions, missing inputs and default values.",
                        ["#", "Type", "Item"], [5, 18, 120])
    for i, (t, txt) in enumerate(items):
        r = i + 4
        _cell(ws8, r, 1, i + 1)
        _cell(ws8, r, 2, t)
        _cell(ws8, r, 3, _txt(txt))
    _add_table(ws8, "Checks", 3, len(items) + 3)

    # ============================== Benchmarks ================================
    ws9 = _titled_sheet(wb, "Benchmarks", "Sanity checks (per sft of covered area)",
                        "Indicative ranges for 5-10 marla houses - CHECK means review the inputs, not that the number is wrong.",
                        ["Check", "Value", "Unit", "Typical Low", "Typical High", "Result"], [40, 12, 10, 11, 11, 10])
    cov_total = sum(f.covered_sft for f in p.floors)
    _cell(ws9, 4, 1, "Total covered area (all floors incl. mumty)", bold=True)
    _cell(ws9, 4, 2, round(cov_total, 1), fmt="#,##0", color=INPUT_FONT_COLOR)
    _cell(ws9, 4, 3, "sft")
    for i, (label, ids, unit, lo, hi) in enumerate(benchmarks_for(p)):
        r = i + 5
        refs = "+".join(f'SUMIF({SCH_ID},"{mid}",{SCH_QTY})' for mid in ids)
        _cell(ws9, r, 1, label)
        _cell(ws9, r, 2, f"=IF($B$4=0,0,({refs})/$B$4)", fmt="0.00")
        _cell(ws9, r, 3, unit)
        _cell(ws9, r, 4, lo, color=INPUT_FONT_COLOR)
        _cell(ws9, r, 5, hi, color=INPUT_FONT_COLOR)
        _cell(ws9, r, 6, f'=IF(B{r}=0,"n/a",IF(AND(B{r}>=D{r},B{r}<=E{r}),"OK","CHECK"))', bold=True)
    rr = len(BENCHMARKS) + 4
    ws9.conditional_formatting.add(f"F5:F{rr}", FormulaRule(formula=['$F5="OK"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws9.conditional_formatting.add(f"F5:F{rr}", FormulaRule(formula=['$F5="CHECK"'], fill=PatternFill("solid", fgColor="FFC7CE")))

    _build_summary(summary, result, mats, sched_row, SCH_ID, SCH_QTY, last)

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------------------
def _build_summary(ws, result: DetailedResult, mats: List[MaterialLine], sched_row, SCH_ID, SCH_QTY, sched_last) -> None:
    p = result.project
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([5, 42, 34, 13, 16, 13, 8, 30, 27, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.cell(row=1, column=2, value="MATERIAL TAKE-OFF  -  WHAT TO BUY").font = _font(size=16, bold=True, color=NAVY)
    ws.cell(row=2, column=2, value=" | ".join(x for x in [p.project_name, p.client, p.location] if x)).font = \
        _font(size=11, bold=True, color=TEAL)
    ws.cell(row=3, column=2, value=f"Scope: {result.scope}   |   Finish: {p.options.finish_tier}   |   "
                                   f"Generated {datetime.now().strftime('%d %b %Y %H:%M')}   |   Quantities only - no costs   |   "
                                   f"App v{_app_version()}, material database {_kb_id()}"
            ).font = _font(italic=True, color="595959")

    r = 4
    if p.drawing_mode == "sketch":
        c = ws.cell(row=r, column=2, value="CONCEPT ESTIMATE from your sketch / requirements (no architectural drawings) - "
                                          "expect about +/-15-30% on the main materials. Re-run with the architect's drawings "
                                          "before final ordering.")
        c.font = _font(size=11, bold=True, color="7F6000")
        c.fill = PatternFill("solid", fgColor="FFE699")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        ws.row_dimensions[r].height = 34
        r += 1
    if p.drawing_mode in ("scanned", "none"):
        msg = ("WARNING: the uploaded drawings could not be read (scanned images/photos). Quantities are for a TYPICAL house "
               "of this plot size, not for these drawings - every line is an assumption." if p.drawing_mode == "scanned" else
               "No drawings were uploaded - quantities are for a TYPICAL house of this plot size (all lines are assumptions).")
        c = ws.cell(row=r, column=2, value=msg)
        c.font = _font(size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="C00000")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        ws.row_dimensions[r].height = 34
        r += 1
    # contents (links filled in at the end, once the shopping-list row is known)
    r += 1
    ws.cell(row=r, column=2, value="CONTENTS  (click to open)").font = _font(size=10, bold=True, color=NAVY)
    contents = [
        ("Shopping list - all materials to buy (further down this page)", None),
        ("Material_Schedule - full detail of every material", "Material_Schedule"),
        ("Procurement_by_Stage - what to buy at each construction stage", "Procurement_by_Stage"),
        ("BOQ_Work_Items - measured quantities", "BOQ_Work_Items"),
        ("Material_Breakdown - how each quantity was calculated", "Material_Breakdown"),
        ("Project_Inputs / Rooms / Openings - data read from the drawings", "Project_Inputs"),
        ("Assumptions_Gaps - things to check before buying", "Assumptions_Gaps"),
        ("Benchmarks - sanity checks", "Benchmarks"),
    ]
    content_rows = []
    for label, target in contents:
        r += 1
        content_rows.append((r, label, target))

    # main materials at a glance
    r += 2
    ws.cell(row=r, column=2, value="MAIN MATERIALS AT A GLANCE").font = _font(size=10, bold=True, color=NAVY)
    r += 1
    for j, h in enumerate(["Material", "Quantity to buy", "Unit"], 2):
        _cell(ws, r, j, h, bold=True, fill=HDR_FILL, color="FFFFFF", align=CENTER)
    for label, ids, unit, div, digits in KEY_CARDS:
        if not any(m.material.mat_id in ids and m.included for m in mats):
            continue
        r += 1
        refs = "+".join(f'SUMIF({SCH_ID},"{mid}",{SCH_QTY})' for mid in ids)
        inner = f"({refs})" if div == 1 else f"({refs})/{div}"
        _cell(ws, r, 2, label, bold=True, fill=CARD_FILL, size=10)
        _cell(ws, r, 3, f"=ROUNDUP({inner},{digits})", bold=True, fill=CARD_FILL,
              fmt="#,##0.00" if digits > 0 else "#,##0", size=11, align=Alignment(horizontal="right", vertical="center"))
        _cell(ws, r, 4, unit, fill=CARD_FILL, size=10)

    # status overview
    r += 2
    ws.cell(row=r, column=2, value="HOW FIRM ARE THE NUMBERS?").font = _font(size=10, bold=True, color=NAVY)
    r += 1
    for j, h in enumerate(["Status", "Lines", "Meaning"], 2):
        _cell(ws, r, j, h, bold=True, fill=HDR_FILL, color="FFFFFF", align=CENTER)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
    meaning = {
        ST_CALC: "Calculated from the drawings / your inputs",
        ST_CALC_ASSUMED: "Calculated, but some inputs are defaults - see Assumptions_Gaps",
        ST_PROVISIONAL: "Allowance for owner-chosen items (lights, AC units)",
        ST_NEEDS_INPUT: "Cannot be calculated yet - an input is missing",
        ST_OPTION: "Optional item not selected (quantity shown if chosen)",
        ST_REFERENCE: "Assembly/duplicate line - already bought under another item",
        ST_NOT_REQ: "Not needed for this house",
    }
    status_rng = f"Material_Schedule!$L$4:$L${sched_last}"
    for st, txt in meaning.items():
        r += 1
        _cell(ws, r, 2, st, fill=PatternFill("solid", fgColor=STATUS_COLORS.get(st, "FFFFFF")))
        _cell(ws, r, 3, f"=COUNTIF({status_rng},B{r})", align=CENTER)
        _cell(ws, r, 4, txt)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
    r += 1
    _cell(ws, r, 2, "Materials in this scope", bold=True)
    _cell(ws, r, 3, f"=COUNTA({SCH_ID})", bold=True, align=CENTER)

    # ------------------------------------------------------ SHOPPING LIST
    r += 3
    shop_anchor = r
    ws.cell(row=r, column=2, value="SHOPPING LIST  -  ALL MATERIALS TO PURCHASE").font = _font(size=13, bold=True, color=NAVY)
    r += 1
    ws.cell(row=r, column=2, value="Grouped by trade. Use the +/- buttons at the left edge to open or close a group. "
                                   "Click an item code for full details.").font = _font(italic=True, color="595959")
    r += 1
    head = ["#", "Material", "Specification / grade", "Quantity to buy", "Buy unit", "Exact qty (incl. wastage)", "Unit",
            "When needed", "Reliability", "Item code"]
    for j, h in enumerate(head, 1):
        _cell(ws, r, j, h, bold=True, fill=HDR_FILL, color="FFFFFF", align=CENTER)
    ws.row_dimensions[r].height = 30
    header_row = r
    cats: Dict[str, List[MaterialLine]] = {}
    for m in mats:
        if m.included and m.gross_qty > 0:
            cats.setdefault(m.material.category, []).append(m)
    n = 0
    for cat, items in cats.items():
        r += 1
        _cell(ws, r, 1, None, fill=CAT_FILL)
        _cell(ws, r, 2, f"{cat}  ({len(items)} items)", bold=True, fill=CAT_FILL, size=10)
        for c in range(3, 11):
            _cell(ws, r, c, None, fill=CAT_FILL)
        for m in items:
            r += 1
            n += 1
            sr = sched_row[m.material.mat_id]
            _, unit = _buy_formula(m, "x")
            st_code = _stage_code(m)
            _cell(ws, r, 1, n, align=Alignment(horizontal="center", vertical="top"))
            _cell(ws, r, 2, m.material.description)
            _cell(ws, r, 3, _txt(m.material.specification))
            _cell(ws, r, 4, f"=Material_Schedule!J{sr}", bold=True, fmt=_buy_fmt(unit), size=10,
                  align=Alignment(horizontal="right", vertical="top"))
            _cell(ws, r, 5, unit)
            _cell(ws, r, 6, f"=Material_Schedule!I{sr}", fmt="#,##0.0", color="595959")
            _cell(ws, r, 7, m.material.unit, color="595959")
            _cell(ws, r, 8, f"{st_code} - {STAGE_NAMES.get(st_code, '')}")
            _cell(ws, r, 9, _friendly_conf(m))
            _link(ws, r, 10, m.material.mat_id, f"'Material_Schedule'!A{sr}", border=True)
            ws.row_dimensions[r].outlineLevel = 1
    shop_last = r
    if shop_last > header_row:
        rng = f"I{header_row + 1}:I{shop_last}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($I{header_row + 1},1)="⚠"'],
                                                       font=Font(name=FONT, size=9, bold=True, color="C65911")))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($I{header_row + 1},1)="✔"'],
                                                       font=Font(name=FONT, size=9, color="548235")))
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.print_title_rows = f"{header_row}:{header_row}"

    pending = [m for m in mats if m.status == ST_NEEDS_INPUT]
    if pending:
        r += 2
        ws.cell(row=r, column=2, value="STILL TO BE QUANTIFIED (an input is missing)").font = _font(size=10, bold=True, color="C00000")
        for m in pending:
            r += 1
            _cell(ws, r, 2, m.material.description)
            _cell(ws, r, 3, _txt(m.material.required_inputs or "quantity"))
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9)
            _link(ws, r, 10, m.material.mat_id, f"'Material_Schedule'!A{sched_row[m.material.mat_id]}", border=True)
    r += 2
    notes = [
        "Quantities include normal site wastage and are rounded UP to how the material is sold "
        "(bags, tons, coils, pipe lengths, hundreds of bricks, tens of cft).",
        "'⚠ Check' items use default values (e.g. window sizes, electrical points) - confirm them before ordering.",
        "Change a wastage % (blue) on Material_Schedule and every quantity on this page updates.",
        "Costs are intentionally excluded. This is a preliminary take-off, not a structural design or a certified QS takeoff.",
    ]
    for txt in notes:
        c = ws.cell(row=r, column=2, value="• " + txt)
        c.font = _font(color="595959")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        ws.row_dimensions[r].height = 24
        r += 1

    for row, label, target in content_rows:
        dest = f"'Summary'!B{shop_anchor}" if target is None else f"'{target}'!A1"
        _link(ws, row, 2, "►  " + label, dest, size=10)
