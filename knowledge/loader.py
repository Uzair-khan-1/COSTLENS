"""
Loader for the Master Material Database (data/master_material_database.xlsx).

The Excel workbook is the single source of truth for materials, work items,
recipes (work item -> material coefficients), engineering coefficients,
nominal mixes, wastage factors and room defaults. Estimators edit the
workbook; the app reads it here.

Coefficient / recipe cells may contain Excel formulas (e.g.
``=Coefficients!$D$12*Mix_Design!$J$4``). They are evaluated in Python by
resolving cell references to Coefficients!D (coefficient values) and
Mix_Design (mix results computed from the mix parts), so the app does NOT
depend on cached Excel values and stays correct after a user edits the
workbook in Excel or LibreOffice.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from functools import lru_cache
from typing import Dict, List, Optional

from openpyxl import load_workbook

DEFAULT_DB_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "master_material_database.xlsx")

_REF = re.compile(r"(?:'?([A-Za-z_][A-Za-z0-9_ ]*)'?!)?\$?([A-Z]{1,2})\$?(\d+)")
_SAFE = re.compile(r"^[0-9eE+\-*/^().\s]*$")


# ---------------------------------------------------------------------------
# Data classes
# ---------------------------------------------------------------------------
@dataclass
class Material:
    mat_id: str
    category: str
    subcategory: str
    description: str
    specification: str
    unit: str
    application: str
    calc_basis: str
    required_inputs: str
    notes: str
    scope: str
    standard: str
    takeoff_unit: str
    formula_key: str
    linked_work_items: str
    wastage_key: str
    stage: str
    tier: str
    in_sample: str
    drawing_source: str
    legacy_code: str
    wastage_pct: float = 0.0  # fraction, e.g. 0.05


@dataclass
class WorkItem:
    wi_id: str
    scope: str
    division: str
    description: str
    unit: str
    unit_si: str
    measurement_rule: str
    quantity_formula: str
    required_inputs: str
    default_spec: str
    legacy_code: str


@dataclass
class RecipeLine:
    wi_id: str
    mat_id: str
    coeff_id: str
    mix_ref: str
    mix_component: str
    coefficient: float
    coeff_unit: str
    basis: str


@dataclass
class Mix:
    mix_id: str
    description: str
    use: str
    cement_parts: float
    sand_parts: float
    agg_parts: float
    dry_factor: float
    values: Dict[str, float] = field(default_factory=dict)  # column letter -> value


@dataclass
class RoomDefault:
    room_type: str
    keywords: List[str]
    floor_finish: str
    skirting: str
    wall_finish: str
    wall_tile_height_ft: float
    ceiling: str
    false_ceiling_pct: float
    points: Dict[str, float]
    notes: str


@dataclass
class KnowledgeBase:
    path: str
    materials: Dict[str, Material]
    work_items: Dict[str, WorkItem]
    recipes: List[RecipeLine]
    coefficients: Dict[str, float]
    coefficient_meta: Dict[str, dict]
    mixes: Dict[str, Mix]
    wastage: Dict[str, float]
    room_defaults: List[RoomDefault]
    scope_map: Dict[str, Dict[str, str]]  # category -> scope -> Y/P/-
    scope_names: List[str]
    stages: Dict[str, str]
    version: str = "1.0"
    file_hash: str = ""  # first 12 hex of SHA-256 of the workbook - printed on every export

    # convenience -----------------------------------------------------------
    def k(self, coeff_id: str) -> float:
        return self.coefficients[coeff_id]

    def recipes_for(self, wi_id: str) -> List[RecipeLine]:
        return [r for r in self.recipes if r.wi_id == wi_id]

    def material_order(self) -> List[str]:
        return list(self.materials.keys())


# ---------------------------------------------------------------------------
# Formula evaluation
# ---------------------------------------------------------------------------
class _Evaluator:
    MIX_COLS = "ABCDEFGHIJKLMNOPQ"

    def __init__(self, coeff_rows: Dict[int, tuple], mix_rows: Dict[int, list]):
        self.coeff_rows = coeff_rows  # row -> (id, raw D value)
        self.mix_rows = mix_rows  # row -> raw row values (A..Q)
        self._coeff_cache: Dict[int, float] = {}
        self._mix_cache: Dict[int, Dict[str, float]] = {}
        self._stack: set = set()

    def coeff(self, row: int) -> float:
        if row in self._coeff_cache:
            return self._coeff_cache[row]
        if row not in self.coeff_rows:
            raise KeyError(f"Coefficients row {row} does not exist")
        if row in self._stack:
            raise ValueError(f"Circular reference at Coefficients!D{row}")
        self._stack.add(row)
        val = self.eval(self.coeff_rows[row][1], "Coefficients")
        self._stack.discard(row)
        self._coeff_cache[row] = val
        return val

    def mix(self, row: int) -> Dict[str, float]:
        if row in self._mix_cache:
            return self._mix_cache[row]
        raw = self.mix_rows[row]
        c, s, a = (float(raw[i] or 0) for i in (3, 4, 5))
        df = self.eval(raw[6], "Mix_Design")
        tot = c + s + a
        v: Dict[str, float] = {"D": c, "E": s, "F": a, "G": df, "H": tot}
        v["I"] = df * c / tot
        bag_cft = self._bag_cft()
        v["J"] = v["I"] / bag_cft
        v["K"] = df * s / tot
        v["L"] = df * a / tot
        v["M"] = v["J"] * self._by_id("K_CFT_M3")
        v["N"] = v["K"]
        v["O"] = v["L"]
        v["P"] = v["J"] * self._by_id("K_BAG_KG") * self._by_id("K_WC")
        self._mix_cache[row] = v
        return v

    def _by_id(self, cid: str) -> float:
        for r, (i, _) in self.coeff_rows.items():
            if i == cid:
                return self.coeff(r)
        raise KeyError(cid)

    def _bag_cft(self) -> float:
        return self._by_id("K_BAG_CFT")

    def eval(self, raw, default_sheet: str) -> float:
        if raw is None or raw == "":
            return 0.0
        if isinstance(raw, (int, float)):
            return float(raw)
        s = str(raw).strip()
        if not s.startswith("="):
            return float(s)
        expr = s[1:]

        def repl(mo):
            sheet = (mo.group(1) or default_sheet).strip()
            col, row = mo.group(2), int(mo.group(3))
            if sheet == "Coefficients" and col == "D":
                return repr(self.coeff(row))
            if sheet == "Mix_Design":
                return repr(self.mix(row)[col])
            raise ValueError(f"Unsupported reference {mo.group(0)} in formula {s}")

        expr2 = _REF.sub(repl, expr)
        if not _SAFE.match(expr2):
            raise ValueError(f"Unsupported formula (only arithmetic allowed): {s}")
        expr2 = expr2.replace("^", "**")
        return float(eval(expr2, {"__builtins__": {}}, {}))  # noqa: S307 - arithmetic only, validated above


# ---------------------------------------------------------------------------
# Loader
# ---------------------------------------------------------------------------
def _rows(ws, min_row=2):
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        if row and row[0] not in (None, ""):
            yield row


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _f(v, default=0.0) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def load_knowledge_base(path: Optional[str] = None) -> KnowledgeBase:
    path = path or DEFAULT_DB_PATH
    return _load_cached(os.path.abspath(path), os.path.getmtime(path))


@lru_cache(maxsize=4)
def _load_cached(path: str, _mtime: float) -> KnowledgeBase:
    wb = load_workbook(path, data_only=False, read_only=False)

    # --- coefficients & mixes (formula aware) ---
    ws = wb["Coefficients"]
    coeff_rows: Dict[int, tuple] = {}
    meta: Dict[str, dict] = {}
    for r in range(2, ws.max_row + 1):
        cid = ws.cell(r, 1).value
        if not cid:
            continue
        coeff_rows[r] = (str(cid).strip(), ws.cell(r, 4).value)
        meta[str(cid).strip()] = {
            "group": _s(ws.cell(r, 2).value), "description": _s(ws.cell(r, 3).value),
            "unit": _s(ws.cell(r, 5).value), "basis": _s(ws.cell(r, 6).value), "source": _s(ws.cell(r, 7).value),
        }
    ws = wb["Mix_Design"]
    mix_rows = {r: [ws.cell(r, c).value for c in range(1, 18)] for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
    ev = _Evaluator(coeff_rows, mix_rows)
    coefficients = {cid: ev.coeff(r) for r, (cid, _) in coeff_rows.items()}
    mixes = {}
    for r, raw in mix_rows.items():
        v = ev.mix(r)
        mixes[str(raw[0])] = Mix(str(raw[0]), _s(raw[1]), _s(raw[2]), v["D"], v["E"], v["F"], v["G"], v)

    # --- wastage ---
    wastage = {str(r[0]).strip(): _f(r[1]) for r in _rows(wb["Wastage_Factors"])}

    # --- materials ---
    materials: Dict[str, Material] = {}
    for r in _rows(wb["Material_Master"]):
        r = list(r) + [None] * 25
        m = Material(
            mat_id=_s(r[0]), category=_s(r[1]), subcategory=_s(r[2]), description=_s(r[3]), specification=_s(r[4]),
            unit=_s(r[5]), application=_s(r[6]), calc_basis=_s(r[7]), required_inputs=_s(r[8]), notes=_s(r[9]),
            scope=_s(r[10]), standard=_s(r[11]), takeoff_unit=_s(r[12]), formula_key=_s(r[13]) or "RECIPE",
            linked_work_items=_s(r[14]), wastage_key=_s(r[15]), stage=_s(r[17]), tier=_s(r[18]), in_sample=_s(r[19]),
            drawing_source=_s(r[20]), legacy_code=_s(r[21]),
        )
        m.wastage_pct = wastage.get(m.wastage_key, 0.0)
        materials[m.mat_id] = m

    # --- work items ---
    work_items: Dict[str, WorkItem] = {}
    for r in _rows(wb["Work_Items"]):
        r = list(r) + [None] * 12
        w = WorkItem(*(_s(x) for x in r[:11]))
        work_items[w.wi_id] = w

    # --- recipes (evaluate coefficient formulas) ---
    recipes: List[RecipeLine] = []
    ws = wb["Recipes"]
    for rr in range(2, ws.max_row + 1):
        wi = ws.cell(rr, 1).value
        if not wi:
            continue
        coef = ev.eval(ws.cell(rr, 8).value, "Recipes")
        recipes.append(RecipeLine(
            wi_id=_s(wi), mat_id=_s(ws.cell(rr, 3).value), coeff_id=_s(ws.cell(rr, 5).value),
            mix_ref=_s(ws.cell(rr, 6).value), mix_component=_s(ws.cell(rr, 7).value), coefficient=coef,
            coeff_unit=_s(ws.cell(rr, 9).value), basis=_s(ws.cell(rr, 10).value),
        ))

    # --- room defaults ---
    ws = wb["Room_Finish_Defaults"]
    hdr = [_s(c.value) for c in ws[1]]
    point_cols = hdr[8:22]
    rooms: List[RoomDefault] = []
    for r in _rows(ws):
        r = list(r)
        rooms.append(RoomDefault(
            room_type=_s(r[0]), keywords=[k.strip().upper() for k in _s(r[1]).split(",") if k.strip()],
            floor_finish=_s(r[2]), skirting=_s(r[3]), wall_finish=_s(r[4]), wall_tile_height_ft=_f(r[5]),
            ceiling=_s(r[6]), false_ceiling_pct=_f(r[7]),
            points={name: _f(v) for name, v in zip(point_cols, r[8:22])}, notes=_s(r[22] if len(r) > 22 else ""),
        ))

    # --- scope map ---
    ws = wb["Scope_Map"]
    hdr = [_s(c.value) for c in ws[1]]
    scope_names = [h for h in hdr[1:] if h and h != "No. of Materials"]
    scope_map: Dict[str, Dict[str, str]] = {}
    for r in _rows(ws):
        cat = _s(r[0])
        if not cat or cat.upper() == "TOTAL" or cat.startswith("Legend"):
            continue
        scope_map[cat] = {name: _s(v) or "-" for name, v in zip(scope_names, r[1:1 + len(scope_names)])}

    stages = {_s(r[0]): _s(r[1]) for r in _rows(wb["Stages"])}
    wb.close()

    import hashlib
    with open(path, "rb") as fh:
        digest = hashlib.sha256(fh.read()).hexdigest()[:12]
    kb = KnowledgeBase(path, materials, work_items, recipes, coefficients, meta, mixes, wastage, rooms,
                       scope_map, scope_names, stages, file_hash=digest)
    validate_knowledge_base(kb)
    return kb


def validate_knowledge_base(kb: KnowledgeBase) -> None:
    """Referential integrity checks - raise ValueError with all problems."""
    problems = []
    for r in kb.recipes:
        if r.mat_id not in kb.materials:
            problems.append(f"Recipe {r.wi_id} -> unknown material {r.mat_id}")
        if r.wi_id not in kb.work_items:
            problems.append(f"Recipe references unknown work item {r.wi_id}")
        if r.coeff_id and r.coeff_id not in kb.coefficients:
            problems.append(f"Recipe {r.wi_id}/{r.mat_id} -> unknown coefficient {r.coeff_id}")
    for m in kb.materials.values():
        if m.wastage_key and m.wastage_key not in kb.wastage:
            problems.append(f"Material {m.mat_id} -> unknown wastage key {m.wastage_key}")
    if problems:
        raise ValueError("Master database integrity errors:\n" + "\n".join(problems[:50]))
