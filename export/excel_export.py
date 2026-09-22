"""
Excel export via openpyxl.

Produces a multi-sheet workbook:
  1. Cover / Project Summary + disclaimer
  2. MTO (with formula/assumption/confidence columns for traceability)
  3. BOQ (with rates, wastage, amounts) - live formulas for wastage
     quantities, amounts, preliminaries, subtotal, contingency and total
  4. Cost Summary (on the Project Summary sheet, linked to the BOQ)
  5. Assumptions & Extraction Notes (everything the AI flagged + defaults used)
"""
from __future__ import annotations

import io
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
from models.schemas import BOQLineItem, CostSummary, ExtractedBuildingParams, ProjectInputs, QuantityLineItem
from utils import units

HEADER_FILL = PatternFill(start_color="1F6FEB", end_color="1F6FEB", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(italic=True, size=10, color="666666")
THIN_BORDER = Border(*(Side(style="thin", color="DDDDDD"),) * 4)

LOW_CONF_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")


def _style_header_row(ws, row_idx: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def _autofit(ws, widths: List[int]):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _plot_label(pi: ProjectInputs) -> str:
    storeys = {1: "single storey", 2: "G+1", 3: "G+2"}.get(pi.plot_storeys, f"{pi.plot_storeys} storeys")
    return f"{pi.plot_marla:g} marla ({pi.marla_sqft:g} sqft/marla), {storeys}"


def _write_cover_sheet(wb: Workbook, project_inputs: ProjectInputs, cost_summary: CostSummary, totals: dict):
    """`totals` holds the BOQ sheet cell addresses of the subtotal,
    contingency and grand total, so the cover figures are LIVE formulas
    that follow any rate/quantity edit made in the BOQ sheet."""
    ws = wb.active
    ws.title = "Project Summary"
    ws["A1"] = config.APP_NAME
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Version {config.APP_VERSION}"
    ws["A2"].font = SUBTITLE_FONT

    cur = cost_summary.currency
    money_fmt = f'"{cur} "#,##0.00'
    rows = [
        ("", ""),
        ("Project Name", project_inputs.project_name),
        ("Client", project_inputs.client_name),
        ("Location", project_inputs.location),
        *([("Plot", _plot_label(project_inputs))] if project_inputs.plot_marla else []),
        ("Soil Type", project_inputs.soil_type),
        ("Concrete Grade (Footing/Column/Beam/Slab)", f"{project_inputs.concrete_grade_footing} / {project_inputs.concrete_grade_column} / {project_inputs.concrete_grade_beam} / {project_inputs.concrete_grade_slab}"),
        ("Steel Grade", project_inputs.steel_grade),
        ("Wall Material", units.relabel_wall_material(project_inputs.wall_material, project_inputs.unit_system)),
        ("Finish Level", project_inputs.finish_level),
        ("Unit System", units.UNIT_SYSTEM_LABELS.get(project_inputs.unit_system, project_inputs.unit_system)),
        ("", ""),
        ("Subtotal", f"=BOQ!{totals['subtotal']}"),
        ("Contingency", f"=BOQ!{totals['contingency']}"),
        ("GRAND TOTAL", f"=BOQ!{totals['grand_total']}"),
    ]
    r = 4
    for label, value in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=label == "GRAND TOTAL")
        cell = ws.cell(row=r, column=2, value=value)
        cell.font = Font(bold=label == "GRAND TOTAL")
        if isinstance(value, str) and value.startswith("=BOQ!"):
            cell.number_format = money_fmt
            cell.alignment = Alignment(horizontal="right")
        r += 1
    ws.cell(row=r, column=1, value="Totals are live formulas linked to the BOQ sheet - edit rates/quantities there.").font = SUBTITLE_FONT
    r += 1

    r += 1
    ws.cell(row=r, column=1, value="DISCLAIMER").font = Font(bold=True, color="C0392B")
    r += 1
    cell = ws.cell(row=r, column=1, value=config.DISCLAIMER_TEXT)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 4, end_column=6)
    _autofit(ws, [38, 30, 14, 14, 14, 14])


def _write_mto_sheet(wb: Workbook, mto_items: List[QuantityLineItem], unit_system: str = units.SI):
    ws = wb.create_sheet("MTO")
    if units.is_fps(unit_system):
        note = (
            f"Quantities shown in {units.UNIT_SYSTEM_LABELS.get(unit_system, unit_system)}. "
            "Formula, Key Inputs Used and Assumptions are all in FPS (ft, in, sqft, cft; steel in kg, cement in bags)."
        )
    else:
        note = (
            f"Quantities shown in {units.UNIT_SYSTEM_LABELS.get(unit_system, unit_system)}. "
            "All calculations are performed internally in SI/metric units; Formula/Key Inputs Used describe that underlying metric arithmetic."
        )
    ws.append([note])
    ws["A1"].font = SUBTITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    headers = ["Item Code", "Category", "Description", "Unit", "Quantity", "Confidence", "Formula", "Key Inputs Used", "Assumptions"]
    ws.append(headers)
    _style_header_row(ws, 2, len(headers))

    for item in mto_items:
        qty, unit = units.quantity_and_unit_for_display(item.quantity, item.unit, unit_system)
        inputs_str = "; ".join(f"{k}={v}" for k, v in item.inputs_used.items())
        assumptions_str = " | ".join(item.assumptions)
        ws.append(
            [
                item.item_code,
                item.category,
                item.description,
                unit,
                round(qty, 3),
                item.confidence.value,
                item.formula,
                inputs_str,
                assumptions_str,
            ]
        )
        row = ws.max_row
        if item.confidence.value == "Low":
            for c in range(1, len(headers) + 1):
                ws.cell(row=row, column=c).fill = LOW_CONF_FILL
        for c in range(1, len(headers) + 1):
            ws.cell(row=row, column=c).border = THIN_BORDER
            ws.cell(row=row, column=c).alignment = Alignment(vertical="top", wrap_text=True)

    _autofit(ws, [14, 14, 34, 8, 12, 12, 42, 40, 50])
    ws.freeze_panes = "A3"


def _write_boq_sheet(
    wb: Workbook,
    boq_items: List[BOQLineItem],
    cost_summary: CostSummary,
    currency_symbol: str,
    unit_system: str = units.SI,
    preliminaries_pct: Optional[float] = None,
) -> dict:
    """Writes the BOQ with LIVE Excel formulas:
       Qty incl. Wastage = Quantity × (1 + Wastage%/100)
       Amount            = Qty incl. Wastage × Rate
       PRELIM-01 rate    = SUM(amounts above) × Preliminaries %
       Subtotal / Contingency / Grand Total = formulas below the table
    so a QS can change any quantity, wastage %, rate or percentage in Excel
    and every total recalculates. Returns the addresses of the total cells."""
    ws = wb.create_sheet("BOQ")
    ws.append([
        f"Quantities/Rates shown in {units.UNIT_SYSTEM_LABELS.get(unit_system, unit_system)}. "
        "Amounts are unaffected by unit system (Quantity x Rate always reproduces the same Amount). "
        "Qty incl. Wastage, Amount and all totals are live formulas."
    ])
    ws["A1"].font = SUBTITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
    headers = ["Item Code", "Category", "Description", "Unit", "Quantity", "Wastage %", "Qty incl. Wastage", "Rate", "Amount", "Confidence", "Remarks"]
    ws.append(headers)
    _style_header_row(ws, 2, len(headers))

    priced_items = [i for i in boq_items if i.item_code != "PRELIM-01"]
    prelim_item = next((i for i in boq_items if i.item_code == "PRELIM-01"), None)
    if preliminaries_pct is None:
        civil = sum(i.amount for i in priced_items)
        preliminaries_pct = (prelim_item.amount / civil * 100.0) if (prelim_item and civil) else 0.0

    first_row = 3
    last_priced_row = first_row + len(priced_items) - 1
    # Parameter cells placed below the table (rows computed up front so the
    # PRELIM-01 formula can reference the % cell).
    n_rows = len(priced_items) + (1 if prelim_item else 0)
    last_data_row = first_row + n_rows - 1
    subtotal_row = last_data_row + 2
    prelim_pct_row = subtotal_row + 1
    cont_pct_row = subtotal_row + 2
    cont_row = subtotal_row + 3
    grand_row = subtotal_row + 4

    qty_fmt = "#,##0.000"
    money_fmt = "#,##0.00"

    def write_row(item: BOQLineItem, row: int, is_prelim: bool = False):
        qty, unit = units.quantity_and_unit_for_display(item.quantity, item.unit, unit_system)
        rate = units.display_rate(item.rate, item.unit, unit_system)
        ws.cell(row=row, column=1, value=item.item_code)
        ws.cell(row=row, column=2, value=item.category)
        ws.cell(row=row, column=3, value=item.description)
        ws.cell(row=row, column=4, value=unit)
        ws.cell(row=row, column=5, value=qty).number_format = qty_fmt
        ws.cell(row=row, column=6, value=item.wastage_pct)
        ws.cell(row=row, column=7, value=f"=E{row}*(1+F{row}/100)").number_format = qty_fmt
        if is_prelim:
            if last_priced_row >= first_row:
                rate_value = f"=SUM(I{first_row}:I{last_priced_row})*$I${prelim_pct_row}/100"
            else:
                rate_value = 0
            ws.cell(row=row, column=8, value=rate_value).number_format = money_fmt
        else:
            ws.cell(row=row, column=8, value=rate).number_format = money_fmt
        ws.cell(row=row, column=9, value=f"=G{row}*H{row}").number_format = money_fmt
        ws.cell(row=row, column=10, value=item.confidence.value)
        ws.cell(row=row, column=11, value=item.remarks)
        fill = item.confidence.value == "Low"
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=c)
            cell.border = THIN_BORDER
            if fill:
                cell.fill = LOW_CONF_FILL

    row = first_row
    for item in priced_items:
        write_row(item, row)
        row += 1
    if prelim_item is not None:
        write_row(prelim_item, row, is_prelim=True)

    bold = Font(bold=True)
    ws.cell(row=subtotal_row, column=8, value="SUBTOTAL").font = bold
    ws.cell(row=subtotal_row, column=9, value=f"=SUM(I{first_row}:I{last_data_row})").font = bold
    ws.cell(row=prelim_pct_row, column=8, value="Preliminaries %")
    ws.cell(row=prelim_pct_row, column=9, value=round(preliminaries_pct, 4))
    ws.cell(row=prelim_pct_row, column=11, value="Input - drives PRELIM-01's rate (% of all items above it).")
    ws.cell(row=cont_pct_row, column=8, value="Contingency %")
    ws.cell(row=cont_pct_row, column=9, value=cost_summary.contingency_pct)
    ws.cell(row=cont_pct_row, column=11, value="Input - % of subtotal.")
    ws.cell(row=cont_row, column=8, value="CONTINGENCY").font = bold
    ws.cell(row=cont_row, column=9, value=f"=I{subtotal_row}*I{cont_pct_row}/100").font = bold
    ws.cell(row=grand_row, column=8, value="GRAND TOTAL").font = bold
    ws.cell(row=grand_row, column=9, value=f"=I{subtotal_row}+I{cont_row}").font = bold
    for r in (subtotal_row, cont_row, grand_row):
        ws.cell(row=r, column=9).number_format = money_fmt
    for r in (prelim_pct_row, cont_pct_row):
        ws.cell(row=r, column=9).font = Font(color="0000FF")  # blue = editable input

    _autofit(ws, [14, 14, 40, 8, 12, 10, 16, 12, 14, 12, 40])
    ws.freeze_panes = "A3"
    return {"subtotal": f"$I${subtotal_row}", "contingency": f"$I${cont_row}", "grand_total": f"$I${grand_row}"}


def _write_assumptions_sheet(wb: Workbook, params: ExtractedBuildingParams):
    ws = wb.create_sheet("Assumptions & Notes")
    ws.append(["AI Overall Notes"])
    ws["A1"].font = Font(bold=True)
    ws.append([params.overall_notes or "-"])
    ws.append([])
    ws.append(["Extraction Warnings"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    if params.extraction_warnings:
        for w in params.extraction_warnings:
            ws.append([f"- {w}"])
    else:
        ws.append(["(none)"])
    _autofit(ws, [110])
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_excel_workbook(
    project_inputs: ProjectInputs,
    params: ExtractedBuildingParams,
    mto_items: List[QuantityLineItem],
    boq_items: List[BOQLineItem],
    cost_summary: CostSummary,
    preliminaries_pct: Optional[float] = None,
) -> bytes:
    unit_system = project_inputs.unit_system
    wb = Workbook()
    _write_mto_sheet(wb, mto_items, unit_system)
    totals = _write_boq_sheet(wb, boq_items, cost_summary, config.DEFAULT_CURRENCY_SYMBOL, unit_system, preliminaries_pct)
    _write_cover_sheet(wb, project_inputs, cost_summary, totals)
    _write_assumptions_sheet(wb, params)
    # Formulas are written without cached values - make Excel compute every
    # formula when the file is opened.
    wb.calculation.fullCalcOnLoad = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
