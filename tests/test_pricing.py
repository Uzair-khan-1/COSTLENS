"""Costing: rate book & provenance, contract types, cash flow, MRS import, government estimate, quotes, price agent."""
from __future__ import annotations

import os
from pathlib import Path

import pytest

from ai.extraction import default_building_params
from detailed_mto import build_project, compute
from engineering import plot_templates
from knowledge import load_knowledge_base
from models.schemas import ProjectInputs
from pricing.base_rates import LABOUR, R
from pricing.costing import CostSettings, compute_cost, pkr, wi_market_costs
from pricing.ratebook import Proposal, RateBook
from scheduling import ScheduleSettings
from scheduling.planner import plan_for_target

ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture(scope="module")
def kb():
    return load_knowledge_base()


@pytest.fixture(scope="module")
def res(kb):
    pi = ProjectInputs(project_name="cost", plot_marla=5)
    params = plot_templates.build_template_params(pi) or default_building_params()
    return compute(build_project(pi, params, kb), kb)


@pytest.fixture()
def book(tmp_path):
    return RateBook(tmp_path)


def test_every_material_and_work_item_has_a_rate(kb):
    assert set(kb.materials) <= set(R), set(kb.materials) - set(R)
    assert set(kb.work_items) <= set(LABOUR)
    for k, (t, lo, hi) in R.items():
        assert lo <= t <= hi, k


def test_rate_lookup_order_and_provenance(book):
    ind = book.rate("CON-001", "Lahore")
    assert ind.status == "indicative" and ind.source
    p = Proposal("Lahore", "CON-001", "Cement", ind.rate, 1500.0, 1480.0, 1520.0, "bag", 0.0,
                 [{"url": "https://www.dawn.com/x", "domain": "dawn.com", "quote": "cement Rs 1,500 per bag", "trusted": True}])
    book.approve(p)
    live = book.rate("CON-001", "Lahore")
    assert live.status == "live" and live.rate == 1500.0 and "dawn" in live.source
    assert book.rate("CON-001", "Faisalabad").rate == 1500.0  # Faisalabad uses Lahore rates
    assert book.rate("CON-001", "Lahore", {"CON-001": 1450.0}).status == "user"
    # followers move with their leader (#3 bar follows the #4 bar)
    p2 = Proposal("Lahore", "RBR-002", "Steel", 262.0, 250.0, 245.0, 255.0, "kg", -4.6, [{"url": "u", "domain": "d"}])
    book.approve(p2)
    assert book.rate("RBR-001", "Lahore").status == "live" and book.rate("RBR-001", "Lahore").follows == "RBR-002"
    book.save()
    again = RateBook(book.folder)
    assert again.rate("CON-001", "Lahore").rate == 1500.0 and again.history


def test_contract_types_and_extras(res, book):
    base = dict(city="Rawalpindi")
    lab = compute_cost(res, CostSettings(contract="labour", **base), book)
    grey = compute_cost(res, CostSettings(contract="grey", **base), book)
    turn = compute_cost(res, CostSettings(contract="turnkey", **base), book)
    assert lab.materials == grey.materials == turn.materials and lab.labour == turn.labour
    assert lab.profit < grey.profit < turn.profit and lab.total < grey.total < turn.total
    assert abs(lab.profit - lab.labour * 0.10) < 1
    assert lab.low < lab.total < lab.high
    ext = compute_cost(res, CostSettings(contract="labour", extras_on={"approval": True, "soil": True}, **base), book)
    assert ext.total > lab.total and len(ext.extras) == 2
    assert 3000 < lab.per_sft < 15000
    assert not any(ln.code == "MSC-008" for ln in lab.lines)  # approvals are an owner extra, not a material


def test_user_quote_changes_the_total(res, book):
    a = compute_cost(res, CostSettings(city="Lahore"), book)
    b = compute_cost(res, CostSettings(city="Lahore", user_rates={"CON-001": 2000.0}), book)
    cem = next(ln for ln in a.lines if ln.code == "CON-001")
    assert b.materials - a.materials == pytest.approx(cem.qty * (2000.0 - cem.rate), rel=1e-6)


def test_cash_flow_follows_the_schedule(res, book):
    plan = plan_for_target(res, 8, ScheduleSettings(start_date="2026-10-05"))
    c = compute_cost(res, CostSettings(escalation_pct_month=1.0), book, plan.schedule)
    assert c.cashflow and c.escalation > 0
    assert abs(sum(m["Total"] for m in c.cashflow) - c.total) / c.total < 0.01
    assert c.cashflow[-1]["Cumulative"] == pytest.approx(sum(m["Total"] for m in c.cashflow))
    flat = compute_cost(res, CostSettings(escalation_pct_month=0.0), book, plan.schedule)
    assert flat.escalation == 0 and flat.total < c.total


def test_rate_books_and_government_estimate(res, book):
    from pricing.government import CITY_BOOK, WI_SPECS, find_item, government_estimate, mrs_index, update_factors
    idx = mrs_index()
    for district, minimum in (("Rawalpindi", 38), ("Lahore", 38), ("Peshawar", 36), ("Karachi", 30)):
        assert district in idx, district
        b = idx[district]
        found = [w for w in WI_SPECS if find_item(b, w)]
        assert len(found) >= minimum, (district, len(found))
    rw = idx["Rawalpindi"]
    assert rw["book"] == "Punjab MRS" and rw["edition"] == "2026-2"
    assert idx["Peshawar"]["book"] == "KP MRS" and idx["Karachi"]["book"] == "Sindh CSR"
    rcc, mult = find_item(rw, "WI-CN-07")
    assert "1: 2: 4" in rcc["description"] and 500 < rcc["composite"] * mult < 1200  # Rs per cft
    c = compute_cost(res, CostSettings(city="Rawalpindi"), book)
    uf = update_factors()
    wc, loose = wi_market_costs(c, res)
    g = government_estimate(res, "Rawalpindi", uf, wc, loose_materials=loose)
    assert g.mapped_share > 0.3 and 0.5 < g.total / c.total < 1.6 and any("BST" in n for n, _v in g.extras)
    # Punjab RCC includes shuttering -> formwork items are not charged twice
    assert all(ln["Amount"] == 0 for ln in g.lines if ln["WI_ID"].startswith("WI-FW-"))
    for city in CITY_BOOK:
        assert government_estimate(res, city, uf, wc, loose_materials=loose).total > 0
    q = government_estimate(res, "Quetta", uf, wc, loose_materials=loose)
    assert any("CSR-2026" in ln["Source"] for ln in q.lines)


def test_older_editions_are_brought_to_today():
    from datetime import date
    from pricing.government import edition_age_years, unit_multiplier
    assert edition_age_years({"edition": "2026-2"}, date(2026, 10, 1)) == 0
    assert 1.3 < edition_age_years({"edition": "2025-1"}, date(2026, 10, 1)) < 1.6
    assert 2.0 < edition_age_years({"edition": "2024"}, date(2026, 10, 1)) < 2.5
    assert unit_multiplier("100 Cft.", "cft") == 0.01 and unit_multiplier("P.Sft", "sft") == 1.0
    assert unit_multiplier("Per Cwt.", "kg") == pytest.approx(1 / 50.8) and unit_multiplier("% Cft", "cft") == 0.01
    assert unit_multiplier("Each", "nos") == 1.0 and unit_multiplier("100 Sft", "cft") is None


@pytest.mark.parametrize("pdf,book,district,edition", [
    ("Market_Rate_System_2nd_Bi_Annual_Rawalpindi_Punjab.pdf", "Punjab MRS", "Rawalpindi", "2026-2"),
    ("Market_Rate_System_2025__1st_Bi_Annual__KPK.pdf", "KP MRS", "Peshawar", "2025-1"),
    ("CSR-2024-_composite_-approved_Sindh.pdf", "Sindh CSR", "Karachi", "2024"),
])
def test_rate_book_parser_on_real_pdfs(pdf, book, district, edition):
    path = f"/mnt/user-data/uploads/{pdf}"
    if not os.path.exists(path):
        pytest.skip("rate book PDF not available")
    from pricing.mrs import parse_mrs
    d = parse_mrs(path)
    assert (d["book"], d["district"], d["edition"]) == (book, district, edition)
    assert len(d["items"]) > 2500


def test_quote_reader_rules_and_units(kb):
    from pricing.quotes import convert, match_items, rule_items
    txt = ("Lucky Cement 50kg OPC   1,460 / bag\nAmreli Sarya G60 12mm   255,000 per ton\n"
           "Bricks A class awwal   17,500 per 1000\nPorcelain floor tiles 24x24   275 sft\nInvoice total 999,999")
    ms = match_items(rule_items(txt), kb.materials.values())
    got = {m.mat_id: m.rate_db_unit for m in ms}
    assert got["CON-001"] == 1460 and got["RBR-002"] == 255.0 and got["MAS-001"] == 17.5 and got["FLR-001"] == 275
    assert len(ms) == 4  # the 'total' line is ignored
    assert convert(40.0, "maund", "kg") is not None and convert(1.0, "sqm", "sft") == pytest.approx(1 / 10.7639)


def test_price_agent_auto_approves_and_queues(book):
    from pricing.agent import PriceAgent
    page = ("<html><body><p>Cement price in Lahore today: Rs 1,470 per bag (Lucky, DG Khan).</p>"
            "<p>Steel sarya rate in Lahore Rs 300 per kg</p></body></html>")
    from pricing.agent import html_to_text

    def search(q):
        return [{"url": "https://www.dawn.com/cement", "title": "rates"}, {"url": "https://unknown-blog.pk/steel", "title": "x"}]

    agent = PriceAgent(book, search=search, fetch=lambda u: html_to_text(page), pause_s=0)
    r1 = agent.update_item("CON-001", "Lahore")
    assert r1.status in ("auto-approved", "pending") and r1.new_rate == 1470
    before = book.rate("CON-001", "Lahore").rate
    if r1.status == "auto-approved":
        assert before == 1470 and book.rate("CON-001", "Lahore").status == "live"
    r2 = agent.update_item("RBR-002", "Lahore")
    assert r2.status in ("pending", "rejected")  # +15 % from the indicative rate: needs a human
    assert any(p.mat_id == "RBR-002" for p in book.pending)


def test_pkr_format():
    assert pkr(15_200_000) == "Rs 1.52 crore" and pkr(4_530_000) == "Rs 45.3 lakh" and pkr(12_500) == "Rs 12,500"


def test_cost_workbook_and_pdf(res, book):
    from pricing.export import build_cost_workbook, cost_pdf
    from openpyxl import load_workbook
    import io
    c = compute_cost(res, CostSettings(), book, plan_for_target(res, 8, ScheduleSettings(start_date="2026-10-05")).schedule)
    wb = load_workbook(io.BytesIO(build_cost_workbook(c, "T")))
    assert {"Cost_Summary", "Materials_Priced", "Labour_Priced", "Cash_Flow"} <= set(wb.sheetnames)
    assert str(wb["Materials_Priced"]["G5"].value).startswith("=E5*F5")
    assert cost_pdf(c, "T")[:4] == b"%PDF"


def test_price_agent_reads_rate_tables_and_skips_old_pages(tmp_path):
    from datetime import date, timedelta
    from pricing.agent import PriceAgent, html_to_text, page_date, table_extract
    d = date.today() - timedelta(days=5)
    old = date.today() - timedelta(days=300)
    pages = {
        "https://a.example/cement": f"<h1>Cement Rate Today {d:%d-%m-%Y}</h1><table><tr><th>Brand</th><th>Price/50 kg bag</th></tr>"
                                    "<tr><td>Lucky Cement</td><td>1,560</td></tr><tr><td>DG Khan Cement</td><td>1,580</td></tr></table>",
        "https://b.example/cement": f"<p>{d:%B %d, %Y}: cement Rs 1,550 per bag in Islamabad</p>",
        "https://c.example/cement": f"<h1>Rates {old:%d %B %Y}</h1><table><tr><td>Lucky Cement</td><td>1,100</td></tr></table>",
    }
    book = RateBook(tmp_path)
    agent = PriceAgent(book, search=lambda q: [], fetch=lambda u: html_to_text(pages[u]), pause_s=0,
                       direct={"CON-001": list(pages)})
    r = agent.update_item("CON-001", "Islamabad")
    assert r.new_rate is not None and 1540 <= r.new_rate <= 1580  # median of current pages; the year is not a price
    assert any("too old" in n for n in r.notes)
    assert {c.domain for c in r.candidates} == {"a.example", "b.example"}  # one vote per website
    assert page_date(f"updated {d:%d-%m-%Y}") == d
    spec = {"sane": (1100, 2200), "keywords": ["cement"], "unit": "bag", "units": {"bag": 1.0}}
    assert table_extract("Cement Rate Today 2026 per bag\nLucky Cement | 1,560 |", spec)[0]["price"] == 1560


def test_news_style_steel_price():
    from pricing.agent import rule_extract
    from pricing.base_rates import LIVE_ITEMS
    got = rule_extract("Grade 60 steel is available at around Rs258 to Rs265 per kilogram.", LIVE_ITEMS["RBR-002"])
    assert got and got[0]["price"] == 261.5
