"""Tests for the Master-Database-driven Detailed MTO (knowledge/ + detailed_mto/)."""
from __future__ import annotations

import os
from io import BytesIO

import pymupdf
import pytest
from openpyxl import load_workbook

from ai.extraction import default_building_params
from detailed_mto import Options, build_project, compute, run_detailed_mto, scan_pdf_bytes
from detailed_mto.edits import apply_openings, openings_to_rows
from detailed_mto.engine import INCLUDED_STATUSES, ST_NEEDS_INPUT, ST_OPTION, ST_SCOPE
from engineering import plot_templates
from knowledge import load_knowledge_base
from models.schemas import ProjectInputs

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLE = os.path.join(ROOT, "sample_data", "sample_5_marla_package.pdf")


@pytest.fixture(scope="module")
def kb():
    return load_knowledge_base()


def _params(pi):
    return plot_templates.build_template_params(pi) or default_building_params(
        wall_thickness_mm=pi.wall_thickness_mm, wall_material=pi.wall_material, unit_system=pi.unit_system)


@pytest.fixture(scope="module")
def pi():
    return ProjectInputs(project_name="Test 5 Marla", plot_marla=5)


# ------------------------------------------------------------------ knowledge base
def test_kb_loads_with_integrity(kb):
    assert len(kb.materials) >= 300
    assert len(kb.work_items) >= 100
    assert len(kb.recipes) >= 300
    for r in kb.recipes:
        assert r.mat_id in kb.materials and r.wi_id in kb.work_items


def test_kb_formulas_evaluated_in_python(kb):
    # bricks per cft from actual brick size 8.75x4.25x2.75 + 3/8 joint
    assert kb.k("K_BRK_PER_CFT") == pytest.approx(13.10, abs=0.01)
    # RCC 1:2:4 nominal mix: ~0.179 bags/cft, ~6.34 bags/m3
    assert kb.mixes["MX_RCC124"].values["J"] == pytest.approx(0.1794, abs=1e-3)
    assert kb.mixes["MX_RCC124"].values["M"] == pytest.approx(6.336, abs=0.01)
    assert kb.mixes["MX_PCC148"].values["J"] == pytest.approx(0.0966, abs=1e-3)
    # recipe coefficient that multiplies a coefficient by a mix cell
    cem = [r for r in kb.recipes if r.wi_id == "WI-CN-07" and r.mat_id == "CON-001"][0]
    assert cem.coefficient == pytest.approx(0.1794, abs=1e-3)


# ------------------------------------------------------------------ engine
def test_every_material_is_listed_with_a_status(kb, pi):
    params = _params(pi)
    result, xlsx = run_detailed_mto(pi, params)
    assert len(result.materials) == len(kb.materials)
    assert all(m.status for m in result.materials)
    ids = [m.material.mat_id for m in result.materials]
    assert ids == list(kb.materials.keys())
    for w in result.work_items:
        assert "error" not in w.calculation.lower(), (w.wi_id, w.calculation)
    wb = load_workbook(BytesIO(xlsx))
    assert wb.sheetnames[:2] == ["Summary", "Material_Schedule"]
    assert wb["Material_Schedule"].max_row >= len(kb.materials) + 3


def test_core_quantities_are_positive_and_consistent(pi):
    result, _ = run_detailed_mto(pi, _params(pi))
    for mid in ("CON-001", "CON-004", "CON-006", "RBR-002", "MAS-001", "FLR-001", "PNT-003", "ELE-006", "PWS-001", "PDR-001"):
        m = result.by_id(mid)
        assert m.included and m.gross_qty > 0, mid
        assert m.gross_qty == pytest.approx(m.net_qty * (1 + m.wastage_pct))
    # cement = sum of its recipe contributions (before wastage)
    cem = result.by_id("CON-001")
    contrib = sum(c.net_qty for c in result.contributions if c.mat_id == "CON-001")
    assert cem.net_qty == pytest.approx(contrib)


def test_scope_filter(pi):
    result, _ = run_detailed_mto(pi, _params(pi), options=Options(scope="Electrical"))
    assert result.by_id("RBR-002").status == ST_SCOPE
    assert result.by_id("ELE-006").status in INCLUDED_STATUSES


def test_options_switch_alternatives(pi):
    base, _ = run_detailed_mto(pi, _params(pi), options=Options(roof_system="Traditional"))
    ins, _ = run_detailed_mto(pi, _params(pi), options=Options(roof_system="Insulated"))
    assert base.by_id("WPF-008").status == ST_OPTION  # insulation board not in traditional roof
    assert ins.by_id("WPF-008").included
    assert not ins.by_id("WPF-009").included  # earth fill only in traditional roof


def test_user_override_changes_quantities(kb, pi):
    params = _params(pi)
    r1, _ = run_detailed_mto(pi, params)
    r2, _ = run_detailed_mto(pi, params, overrides={"N_WC": 6})
    assert r2.by_id("SAN-001").gross_qty == 6
    assert r1.by_id("SAN-001").gross_qty != 6


def test_opening_edits(kb, pi):
    p = build_project(pi, _params(pi), kb)
    rows = openings_to_rows(p)
    before = sum(o.area_total for o in p.windows())
    for r in rows:
        if r["Kind"] == "window":
            r["Width (ft)"] = r["Width (ft)"] + 1
    apply_openings(p, rows, openings_to_rows(build_project(pi, _params(pi), kb)))
    assert sum(o.area_total for o in p.windows()) > before
    assert any(o.confidence == "User" for o in p.windows())
    res = compute(p, kb)
    assert res.by_id("DWG-009").gross_qty > 0


def test_existing_app_objects_not_mutated(pi):
    params = _params(pi)
    before = params.model_dump_json()
    pi_before = pi.model_dump_json()
    run_detailed_mto(pi, params, overrides={"H_FLOOR": 13})
    assert params.model_dump_json() == before
    assert pi.model_dump_json() == pi_before


# ------------------------------------------------------------------ scanner
def _synthetic_pdf() -> bytes:
    doc = pymupdf.open()
    pg = doc.new_page()
    y = 60
    for t in ["GROUND FLOOR PLAN", "PLUMBING WORK DETAILS", "F.T", "F.T", "F.T", "M.H", "VANITY", "CONCEALED WC", "wc",
              "SHOWER AREA", "SEPTIC TANK", "8'x4'"]:
        pg.insert_text((50, y), t, fontsize=9)
        y += 14
    pg2 = doc.new_page()
    pg2.insert_text((60, 60), "FULL HOUSE CHOGATTH", fontsize=9)
    pg2.insert_text((60, 100), "3'-6\" X 7'-0\"", fontsize=9)
    pg2.insert_text((200, 100), "SINGLE", fontsize=9)
    pg2.insert_text((300, 100), "6", fontsize=9)
    pg2.insert_text((360, 100), "ROOM DOOR", fontsize=9)
    pg2.insert_text((60, 130), "4'-6\" X 8'-0\"   10\"", fontsize=9)
    pg2.insert_text((200, 130), "DOUBLE", fontsize=9)
    pg2.insert_text((300, 130), "1", fontsize=9)
    pg2.insert_text((360, 130), "MAIN DOOR", fontsize=9)
    pg3 = doc.new_page()
    pg3.insert_text((60, 60), "MUMTY PLAN", fontsize=9)
    pg3.insert_text((60, 80), "PLUMBING WORK DETAILS", fontsize=9)
    pg3.insert_text((60, 100), "600 GAL. FG WATER TANK", fontsize=9)
    return doc.tobytes()


def test_scanner_reads_labels_and_door_schedule():
    res = scan_pdf_bytes([{"name": "synthetic.pdf", "bytes": _synthetic_pdf()}])
    assert res.get("FT") == 3
    assert res.get("MH") == 1
    assert res.get("WC") == 1  # 'CONCEALED WC' + 'wc' tag on one fixture counts once
    assert res.oh_tank_gal == 600
    doors = {d.name: d for d in res.doors}
    assert doors["Room Door"].qty == 6 and doors["Room Door"].width_ft == 3.5
    assert doors["Main Door"].leaves == 2 and doors["Main Door"].chogath_in == 10


def test_scanner_never_raises_on_bad_input():
    res = scan_pdf_bytes([{"name": "x.pdf", "bytes": b"not a pdf"}, {"name": "img.png", "bytes": b"123"}])
    assert res.label_counts == {}


@pytest.mark.skipif(not os.path.exists(SAMPLE), reason="sample package not present")
def test_sample_package_end_to_end(pi):
    from drawing_processing.package_analyzer import analyze_package, apply_package_facts
    files = [{"name": "sample_5_marla_package.pdf", "bytes": open(SAMPLE, "rb").read(), "view_tag": "Auto"}]
    facts = analyze_package(files, None, "FPS")
    params, _, _ = apply_package_facts(_params(pi), facts, pi)
    result, xlsx = run_detailed_mto(pi, params, facts=facts, files=files)
    counts = result.status_counts()
    assert sum(v for k, v in counts.items() if k in INCLUDED_STATUSES) > 150
    assert counts.get(ST_NEEDS_INPUT, 0) < 20
    assert len(xlsx) > 10000


# ------------------------------------------------------------------ v0.5 refinements
def test_rcc_mix_option_changes_structural_cement_only(pi):
    params = _params(pi)
    a, _ = run_detailed_mto(pi, params, options=Options(rcc_mix="MX_RCC124"))
    b, _ = run_detailed_mto(pi, params, options=Options(rcc_mix="MX_RCC1153"))
    slab_a = [c for c in a.contributions if c.wi_id == "WI-CN-07" and c.mat_id == "CON-001"][0]
    slab_b = [c for c in b.contributions if c.wi_id == "WI-CN-07" and c.mat_id == "CON-001"][0]
    assert slab_b.coefficient > slab_a.coefficient * 1.2  # 1:1.5:3 is richer
    tank_a = [c for c in a.contributions if c.wi_id == "WI-CN-11" and c.mat_id == "CON-001"][0]
    tank_b = [c for c in b.contributions if c.wi_id == "WI-CN-11" and c.mat_id == "CON-001"][0]
    assert tank_a.coefficient == pytest.approx(tank_b.coefficient)  # tanks keep their drawing mix


def test_rooms_override_drives_derived_counts(kb, pi):
    from detailed_mto.model import Room
    rooms = [Room("ground", "BATH 1", "Bathroom", 8, 5), Room("ground", "BATH 2", "Bathroom", 8, 5),
             Room("ground", "BATH 3", "Bathroom", 8, 5), Room("ground", "KITCHEN", "Kitchen", 10, 9),
             Room("ground", "BED", "Bedroom", 12, 12)]
    p = build_project(pi, _params(pi), kb, rooms_override=rooms)
    assert p.v("N_BATH") == 3 and p.v("N_WC") == 3 and p.v("N_KSINK") == 1
    assert len(p.rooms) == 5


def test_export_has_no_cost_columns(pi):
    _, xlsx = run_detailed_mto(pi, _params(pi))
    wb = load_workbook(BytesIO(xlsx))
    headers = [c.value for c in wb["Material_Schedule"][3]]
    assert not any("Rate" in (h or "") or "Amount" in (h or "") for h in headers)


def test_row_conversions_roundtrip(kb, pi):
    from detailed_mto.edits import mark_user_edits, rooms_to_rows, rows_to_rooms
    p = build_project(pi, _params(pi), kb)
    rows = rooms_to_rows(p)
    assert len(rows_to_rooms(rows)) == len([r for r in p.rooms if r.area > 0])
    changed = [dict(r) for r in rows]
    changed[0]["Length (ft)"] = changed[0]["Length (ft)"] + 2
    marked = mark_user_edits(changed, rows, ["Floor", "Room", "Room type", "Length (ft)", "Width (ft)"])
    assert marked[0]["Confidence"] == "User" and marked[1]["Confidence"] == rows[1]["Confidence"]


def test_benchmarks_depend_on_structural_system(kb, pi):
    from detailed_mto.export import benchmarks_for
    p = build_project(pi, _params(pi), kb)
    steel = [b for b in benchmarks_for(p) if b[2] == "kg/sft"][0]
    assert ("RCC frame" in steel[0]) == (p.v("FDN_STRIP") < 0.5)


# ------------------------------------------------------------------ v0.5.1 export & scope
def test_export_excludes_out_of_scope_materials(pi):
    res, xlsx = run_detailed_mto(pi, _params(pi), options=Options(scope="Electrical"))
    wb = load_workbook(BytesIO(xlsx))
    ids = [c.value for c in wb["Material_Schedule"]["A"][3:] if c.value]
    assert ids and all(i in {m.material.mat_id for m in res.scoped_materials()} for i in ids)
    assert "RBR-002" not in ids and "CON-001" not in ids
    assert not any(m.status == ST_SCOPE for m in res.scoped_materials())
    cats = {c.value for c in wb["Material_Schedule"]["B"][3:] if c.value}
    assert "Reinforcement & Steel" not in cats


def test_summary_has_shopping_list_with_every_purchased_material(pi):
    res, xlsx = run_detailed_mto(pi, _params(pi))
    ws = load_workbook(BytesIO(xlsx))["Summary"]
    codes = {c.value for c in ws["J"] if isinstance(c.value, str) and "-" in c.value}
    assert {m.material.mat_id for m in res.purchase_list()} <= codes
    texts = [c.value for c in ws["B"] if isinstance(c.value, str)]
    assert any("SHOPPING LIST" in t for t in texts)
    assert any(ws.cell(row=r, column=2).hyperlink is not None for r in range(5, 10))  # contents links


def test_purchase_units(kb):
    from detailed_mto.engine import purchase_qty
    assert purchase_qty(kb.materials["CON-001"], 954.2) == (955, "bags (50 kg)")
    assert purchase_qty(kb.materials["RBR-002"], 3233.1) == (3.24, "ton")
    assert purchase_qty(kb.materials["ELE-006"], 825) == (10, "coils (90 m)")
    assert purchase_qty(kb.materials["MAS-001"], 83594.9)[0] == 83600


# ------------------------------------------------------------------ v0.5.2 fixes
def test_overrides_propagate_to_derived_values(kb, pi):
    params = _params(pi)
    p0 = build_project(pi, params, kb)
    p1 = build_project(pi, params, kb, overrides={"H_FLOOR": p0.v("H_FLOOR") + 2, "PLOT_W": 60})
    assert all(f.storey_height_ft == p1.v("H_FLOOR") for f in p1.storeys)
    assert p1.v("H_TOTAL") > p0.v("H_TOTAL")
    assert p1.v("SEWER_LEN") == 70
    assert p1.params["H_FLOOR"].confidence == "User"
    r0, r1 = compute(p0, kb), compute(p1, kb)
    assert r1.by_id("MAS-001").gross_qty > r0.by_id("MAS-001").gross_qty  # taller walls -> more bricks
    res, _ = run_detailed_mto(pi, params, overrides={"H_FLOOR": 13})
    assert res.project.floors[0].storey_height_ft == 13


def test_validation_flags_impossible_and_unusual_values(kb, pi):
    from detailed_mto.validation import errors, validate_project
    params = _params(pi)
    assert not errors(validate_project(build_project(pi, params, kb)))
    bad = build_project(pi, params, kb, overrides={"N_WC": -2, "H_FLOOR": 0, "T_SLAB_IN": 30})
    keys = {i.key for i in errors(validate_project(bad))}
    assert {"N_WC", "H_FLOOR", "T_SLAB_IN"} <= keys
    odd = build_project(pi, params, kb, overrides={"ROOF_AREA": 9000})
    sev = {i.key: i.severity for i in validate_project(odd)}
    assert sev.get("ROOF_AREA") == "warning"


def test_negative_inputs_never_create_negative_quantities(kb, pi):
    res, _ = run_detailed_mto(pi, _params(pi), overrides={"N_WC": -5, "SEWER_LEN": -100})
    assert all(m.gross_qty >= 0 for m in res.materials)


def test_scanned_drawings_are_marked_assumed(kb, pi):
    files = [{"name": "photo.png", "bytes": b"\x89PNG....", "view_tag": "Plan"}]
    res, xlsx = run_detailed_mto(pi, _params(pi), files=files)
    p = res.project
    assert p.drawing_mode == "scanned"
    assert all(prm.confidence == "Assumed" for prm in p.params.values())
    assert res.status_counts().get("Calculated", 0) == 0  # nothing claims to come from the drawings
    assert any("could not be read" in c or "scanned" in c for c in p.conflicts)
    ws = load_workbook(BytesIO(xlsx))["Summary"]
    assert "WARNING" in (ws["B4"].value or "")


def test_no_drawings_mode(kb, pi):
    res, _ = run_detailed_mto(pi, _params(pi))
    assert res.project.drawing_mode == "none"
    assert res.status_counts().get("Calculated", 0) == 0


def test_unedited_rows_are_not_marked_as_user_edits(kb, pi):
    import pandas as pd
    from detailed_mto.edits import OPENING_COLS, ROOM_COLS, mark_user_edits, openings_to_rows, rooms_to_rows
    p = build_project(pi, _params(pi), kb)
    for rows, cols, keys in ((rooms_to_rows(p), ROOM_COLS, ["Floor", "Room", "Room type", "Length (ft)", "Width (ft)"]),
                             (openings_to_rows(p), OPENING_COLS, ["Kind", "Name", "Width (ft)", "Height (ft)", "Qty", "Leaves", "External"])):
        via_editor = pd.DataFrame(rows, columns=cols).to_dict("records")  # ints become floats, like st.data_editor
        marked = mark_user_edits(via_editor, rows, keys)
        assert not any(r["Confidence"] == "User" for r in marked)


def test_editor_delta_applied_when_leaving_step_3():
    from ui.mto_views import _apply_editor_delta
    rows = [{"Room": "A", "Length (ft)": 10}, {"Room": "B", "Length (ft)": 12}, {"Room": "C", "Length (ft)": 8}]
    delta = {"edited_rows": {0: {"Length (ft)": 11}}, "deleted_rows": [1], "added_rows": [{"Room": "D", "Length (ft)": 9}]}
    out = _apply_editor_delta(rows, delta, ["Room", "Length (ft)"])
    assert [r["Room"] for r in out] == ["A", "C", "D"] and out[0]["Length (ft)"] == 11
