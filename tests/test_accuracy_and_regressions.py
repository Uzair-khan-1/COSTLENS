"""
Range / benchmark and regression tests.

Unlike tests/test_calculations.py (which locks formulas in place), these
check that the outputs are REALISTIC against standard estimating norms and
published market benchmarks, and guard the specific bugs fixed in v0.4:
footing depth vs thickness, brick-count double deduction, beam/slab
overlap, plaster split, AI enum/material handling, validation, grade
adjustments, Groq model fallback and live Excel formulas.
"""
import io
import math

import pytest
from openpyxl import load_workbook

from ai.extraction import (
    _map_json_to_params,
    _strip_code_fences,
    default_building_params,
    normalize_footing_type,
    resolve_wall_material,
)
from engineering import rules
from engineering.calculations import (
    compute_beam_concrete,
    compute_footing_concrete,
    compute_masonry,
    compute_wall_face_areas,
    concrete_material_breakdown,
    generate_all_quantities,
)
from engineering.validation import validate_params
from export.excel_export import build_excel_workbook
from models.schemas import ConfidenceLevel, EngineeringAssumptions, Estimate, ProjectInputs, Source, WastageFactors
from mto_boq.boq_generator import boq_to_dataframe, concrete_grade_rate_delta, generate_boq
from mto_boq.rates import load_default_rates

SQFT_PER_SQM = 10.7639


def E(v, c=ConfidenceLevel.HIGH):
    return Estimate(value=v, confidence=c, source=Source.USER_INPUT)


def rate_book():
    return {r.item_code: r for r in load_default_rates()}


def five_marla_g1(material="Burnt clay brick (traditional 230x110x75mm)"):
    p = default_building_params(wall_material=material)
    p.num_floors = E(2)
    p.plinth_area_per_floor_sqm = E(93)
    p.slabs.area_per_floor_sqm = E(93)
    p.footings.count = E(12)
    p.footings.length_m = E(1.5)
    p.footings.width_m = E(1.5)
    p.columns.count = E(12)
    p.beams.count = E(16)
    p.walls.total_length_per_floor_m = E(75)
    p.walls.external_perimeter_m = E(39)
    p.openings.door_count_per_floor = E(7)
    p.openings.window_count_per_floor = E(8)
    p.services.bathroom_count_total = E(4)
    return p, ProjectInputs(wall_material=material)


def by_code(items):
    return {i.item_code: i for i in items}


# --------------------------------------------------------------------------
# Standard estimating norms
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "material, lo, hi",
    [
        ("Burnt clay brick (modular 190x90x90mm)", 490, 510),  # textbook: 500 bricks/m3
        ("Burnt clay brick (traditional 230x110x75mm)", 395, 420),  # 9"x4.5"x3" brick
    ],
)
def test_bricks_per_m3_match_standard_norms(material, lo, hi):
    walls = default_building_params(wall_material=material).walls
    walls.total_length_per_floor_m = E(10)
    walls.height_m = E(1)
    walls.thickness_m = E(0.1)  # 1 m3 of wall
    openings = default_building_params().openings
    for f in ("door_count_per_floor", "window_count_per_floor"):
        setattr(openings, f, E(0))
    items, _ = compute_masonry(walls, openings, num_floors=1)
    units = by_code(items)["MAS-01"].quantity
    mortar = by_code(items)["MAS-02"].quantity
    assert lo <= units <= hi
    assert 0.20 <= mortar <= 0.26  # wet mortar per m3 of brickwork


def test_m20_cement_bags_per_m3():
    bags = concrete_material_breakdown(1.0, "M20")["cement_bags"]
    assert math.isclose(bags, 8.06, abs_tol=0.05)


def test_footing_concrete_uses_thickness_not_founding_depth():
    p = default_building_params()
    item = compute_footing_concrete(p.footings, "M20")
    assert math.isclose(item.quantity, 9 * 1.2 * 1.2 * rules.DEFAULT_FOOTING_THICKNESS_M, rel_tol=1e-3)


def test_excavation_deeper_than_footing_and_backfill_positive():
    p, pi = five_marla_g1()
    items = by_code(generate_all_quantities(p, pi))
    assert items["EXC-01"].quantity > 3 * items["FTG-CONC-01"].quantity
    assert items["BACKFILL-01"].quantity > 0


def test_beam_excludes_slab_overlap():
    p = default_building_params()
    full = compute_beam_concrete(p.beams, 2, "M20", slab_thickness_m=0.0).quantity
    net = compute_beam_concrete(p.beams, 2, "M20", slab_thickness_m=0.125).quantity
    b = p.beams
    expected = b.count.value * b.avg_length_m.value * b.width_m.value * (0.45 + 2 * (0.45 - 0.125))
    assert net < full
    assert math.isclose(net, round(expected, 3), rel_tol=1e-6)


def test_plaster_faces_add_up_to_both_faces_of_all_walls():
    p, pi = five_marla_g1()
    _, net = compute_masonry(p.walls, p.openings, 2)
    fa = compute_wall_face_areas(p, net, 2)  # no parapet
    assert math.isclose(fa["internal_sqm"] + fa["external_sqm"], 2 * net, rel_tol=1e-9)
    assert fa["internal_sqm"] > fa["external_sqm"]  # internal walls are plastered both sides


def test_steel_ratio_in_typical_range():
    p, pi = five_marla_g1()
    items = generate_all_quantities(p, pi)
    conc = sum(i.quantity for i in items if i.category in {"Footing", "Column", "Beam", "Slab", "Staircase", "Lintels"} and i.unit == "m3" and not i.informational)
    steel = sum(i.quantity for i in items if i.category == "Reinforcement")
    assert 75 <= steel / conc <= 160


@pytest.mark.parametrize("scenario", ["default", "five_marla"])
def test_cost_per_sqft_within_published_turnkey_band(scenario):
    """Published 2026 Pakistani turnkey (grey + standard finishing) ranges
    cluster around PKR 5,800-8,800 per sq ft; allow a margin either side."""
    if scenario == "default":
        p, pi = default_building_params(), ProjectInputs()
    else:
        p, pi = five_marla_g1()
    mto = generate_all_quantities(p, pi)
    _, cs = generate_boq(mto, rate_book(), WastageFactors(), pi)
    sqft = p.plinth_area_per_floor_sqm.value * p.num_floors.value * SQFT_PER_SQM
    assert 5000 <= cs.grand_total / sqft <= 9500


def test_every_priced_item_has_a_rate():
    p, pi = five_marla_g1()
    boq, _ = generate_boq(generate_all_quantities(p, pi), rate_book(), WastageFactors(), pi)
    missing = [b.item_code for b in boq if "No rate found" in b.remarks]
    assert missing == []


def test_fps_display_invariance():
    p, pi = five_marla_g1()
    boq, _ = generate_boq(generate_all_quantities(p, pi), rate_book(), WastageFactors(), pi)
    si, fps = boq_to_dataframe(boq, "SI"), boq_to_dataframe(boq, "FPS")
    rel = abs((si.Quantity * si.Rate).sum() - (fps.Quantity * fps.Rate).sum()) / (si.Quantity * si.Rate).sum()
    assert rel < 1e-4


def test_assumption_defaults_mirror_rules():
    A = EngineeringAssumptions()
    for member in ("footing", "column", "beam", "slab", "stair", "lintel"):
        assert getattr(A, f"steel_kg_per_m3_{member}") == rules.STEEL_THUMB_RULE_KG_PER_M3[member][1]
    assert A.pcc_thickness_m == rules.DEFAULT_PCC_THICKNESS_M
    assert A.plinth_height_m == rules.DEFAULT_PLINTH_HEIGHT_M
    assert A.excavation_working_space_m == rules.EXCAVATION_WORKING_SPACE_M
    assert A.parapet_height_m == rules.DEFAULT_PARAPET_HEIGHT_M


def test_editable_steel_thumb_rule_is_used():
    p, pi = five_marla_g1()
    base = by_code(generate_all_quantities(p, pi))["SLAB-STEEL-01"].quantity
    more = by_code(generate_all_quantities(p, pi, EngineeringAssumptions(steel_kg_per_m3_slab=100)))["SLAB-STEEL-01"].quantity
    assert math.isclose(more / base, 100 / 85, rel_tol=1e-3)


# --------------------------------------------------------------------------
# Grades
# --------------------------------------------------------------------------


def test_steel_grade_factor_and_alignment():
    assert rules.steel_grade_qty_factor("Fe415") == rules.steel_grade_qty_factor("Grade 60 (60,000 psi)") == 1.0
    assert rules.steel_grade_qty_factor("Grade 40 (40,000 psi)") > 1.0
    assert rules.steel_grade_qty_factor("Fe500") < 1.0


def test_concrete_grade_rate_delta():
    assert concrete_grade_rate_delta("M20", "M20") == 0
    assert concrete_grade_rate_delta("M25", "M20") > 0
    assert concrete_grade_rate_delta("M15", "M20") < 0
    assert math.isclose(concrete_grade_rate_delta("3600 psi", "M20"), concrete_grade_rate_delta("M25", "M20"))


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------


def test_defaults_pass_validation_cleanly():
    errors, warnings = validate_params(default_building_params(), ProjectInputs())
    assert errors == [] and warnings == []


def test_validation_blocks_impossible_inputs():
    p = default_building_params()
    p.openings.door_count_per_floor = E(200)
    p.num_floors = E(0)
    p.footings.founding_depth_m = E(0.3)
    errors, _ = validate_params(p)
    joined = " ".join(errors)
    assert "at least 1" in joined and "larger than the total wall area" in joined and "Founding depth" in joined


# --------------------------------------------------------------------------
# AI output handling
# --------------------------------------------------------------------------


def test_footing_type_normalization():
    assert normalize_footing_type("Isolated") == "isolated"
    assert normalize_footing_type("Pad footing") == "isolated"
    assert normalize_footing_type("MAT") == "raft"
    assert normalize_footing_type(None) == "isolated"


def test_wall_material_never_overrides_step1_with_free_text():
    step1 = "Burnt clay brick (traditional 230x110x75mm)"
    assert resolve_wall_material("brick", step1) == step1
    assert resolve_wall_material("9 inch brick masonry", step1) == step1
    assert resolve_wall_material("AAC blocks", step1).startswith("AAC")


def test_partial_ai_json_keeps_good_fields_and_defaults_the_rest():
    raw = {"num_floors": ["2", "H"], "footings": {"footing_type": "Pad", "count": [12, "H"], "depth_m": "0.5 m"}}
    p = _map_json_to_params(raw, 230, "Burnt clay brick (modular 190x90x90mm)")
    assert p.num_floors.value == 2 and p.footings.count.value == 12 and p.footings.depth_m.value == 0.5
    assert p.footings.footing_type == "isolated"
    assert p.columns.count.confidence == ConfidenceLevel.LOW  # defaulted
    assert any("did not return" in w for w in p.extraction_warnings)


def test_strip_code_fences():
    assert _strip_code_fences('```json\n{"a": 1}\n```') == '{"a": 1}'


# --------------------------------------------------------------------------
# Groq client reliability
# --------------------------------------------------------------------------


def test_groq_client_caps_images_retries_and_falls_back(monkeypatch):
    from PIL import Image

    import ai.groq_client as g
    import config

    monkeypatch.setattr(config, "GROQ_RETRY_BASE_DELAY_S", 0)
    calls = []

    class Err(Exception):
        def __init__(self, code, msg):
            super().__init__(msg)
            self.status_code = code

    script = [Err(400, "reasoning_effort unsupported"), Err(429, "rate limit"), Err(404, "model decommissioned"), "OK"]

    class Resp:
        def __init__(self, text):
            self.choices = [type("C", (), {"message": type("M", (), {"content": text})()})()]

    class Client:
        class chat:
            class completions:
                @staticmethod
                def create(**kw):
                    calls.append((kw["model"], "reasoning_effort" in kw, len(kw["messages"][1]["content"]) - 1))
                    step = script.pop(0)
                    if isinstance(step, Exception):
                        raise step
                    return Resp(step)

    monkeypatch.setattr(g, "get_client", lambda key: Client)
    out = g.call_vision_model("k", "s", "u", [Image.new("RGB", (8, 8))] * 6)
    assert out == "OK"
    assert all(n_images == config.MAX_IMAGES_PER_REQUEST for _, _, n_images in calls)
    assert calls[0][0] == config.GROQ_VISION_MODEL and calls[-1][0] == config.GROQ_VISION_MODEL_FALLBACK


# --------------------------------------------------------------------------
# Excel export
# --------------------------------------------------------------------------


def test_excel_boq_uses_live_formulas():
    p, pi = five_marla_g1()
    mto = generate_all_quantities(p, pi)
    boq, cs = generate_boq(mto, rate_book(), WastageFactors(), pi)
    wb = load_workbook(io.BytesIO(build_excel_workbook(pi, p, mto, boq, cs, preliminaries_pct=8.0)))
    ws = wb["BOQ"]
    assert str(ws["G3"].value).startswith("=E3*")
    assert str(ws["I3"].value) == "=G3*H3"
    assert str(wb["Project Summary"]["B17"].value).startswith("=BOQ!")
