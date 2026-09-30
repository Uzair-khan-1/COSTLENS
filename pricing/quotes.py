"""
Supplier quotation reader.

The owner uploads a dealer's quotation or bill (PDF, photo, Excel/CSV or pasted text). Items and prices are
extracted - by the AI when a key is available (it can read photos), otherwise by rules from the text - then
matched to the materials of the project and converted to the database unit (per bag, per kg, per brick ...).
The owner confirms each match; confirmed prices become project-specific rates ("Your quote").
"""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from typing import List, Optional, Tuple

SYNONYMS = {
    "sarya": "steel deformed bar", "saria": "steel deformed bar", "rebar": "deformed bar", "g60": "grade 60 deformed bar",
    "bajri": "crushed stone crush", "rora": "crushed stone crush", "crush": "crushed stone crush", "rait": "sand", "ret": "sand",
    "eent": "brick", "int": "brick", "bricks": "brick", "awwal": "class a brick", "cement": "portland cement",
    "wire": "copper wire", "3/29": "3/29 wire", "7/29": "7/29 wire", "7/36": "7/36 wire", "ppr": "ppr pipe", "upvc": "upvc pipe",
    "tiles": "tiles", "tile": "tiles", "emulsion": "emulsion paint", "weathershield": "weather exterior emulsion",
}

# quote unit -> (multiplier to reach a common base, base) ; material units map to the same bases
UNIT_BASE = {
    "bag": (1, "bag"), "bags": (1, "bag"), "kg": (1, "kg"), "ton": (1000, "kg"), "tonne": (1000, "kg"), "mt": (1000, "kg"),
    "maund": (40, "kg"), "mann": (40, "kg"), "brick": (1, "nos"), "nos": (1, "nos"), "no": (1, "nos"), "piece": (1, "nos"),
    "pcs": (1, "nos"), "each": (1, "nos"), "set": (1, "set"), "1000": (1000, "nos"), "thousand": (1000, "nos"),
    "cft": (1, "cft"), "100 cft": (100, "cft"), "sft": (1, "sft"), "sq ft": (1, "sft"), "sqft": (1, "sft"),
    "sqm": (10.7639, "sft"), "m2": (10.7639, "sft"), "rft": (1, "rft"), "ft": (1, "rft"), "foot": (1, "rft"),
    "m": (1, "m"), "meter": (1, "m"), "metre": (1, "m"), "coil": (90, "m"), "roll": (90, "m"), "gallon": (1, "gal"),
    "gal": (1, "gal"), "litre": (1 / 3.64, "gal"), "liter": (1 / 3.64, "gal"), "l": (1 / 3.64, "gal"), "ls": (1, "ls"),
    "length": (13.12, "rft"),
}
DB_BASE = {"bag": (1, "bag"), "kg": (1, "kg"), "Nos": (1, "nos"), "nos": (1, "nos"), "set": (1, "set"), "cft": (1, "cft"),
           "sft": (1, "sft"), "rft": (1, "rft"), "m": (1, "m"), "gal": (1, "gal"), "LS": (1, "ls"), "L": (1 / 3.64, "gal"),
           "ton": (1000, "kg")}


@dataclass
class QuoteItem:
    text: str
    price: float
    unit: str
    qty: Optional[float] = None


@dataclass
class QuoteMatch:
    item: QuoteItem
    mat_id: Optional[str]
    material: str
    score: float
    rate_db_unit: Optional[float]  # price converted to the material's database unit
    db_unit: str
    note: str = ""


# ---------------------------------------------------------------------------
# reading
# ---------------------------------------------------------------------------
def file_text(name: str, data: bytes) -> str:
    n = name.lower()
    if n.endswith(".pdf"):
        import pymupdf
        doc = pymupdf.open(stream=data, filetype="pdf")
        return "\n".join(p.get_text() for p in list(doc)[:10])
    if n.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets[:3]:
            for r in ws.iter_rows(values_only=True):
                cells = [str(c) for c in r if c not in (None, "")]
                if cells:
                    rows.append("  ".join(cells))
        return "\n".join(rows[:800])
    if n.endswith((".csv", ".txt")):
        return data.decode("utf-8", errors="replace")[:60000]
    return ""


_UNIT_WORDS = "|".join(sorted((re.escape(u) for u in UNIT_BASE), key=len, reverse=True))
_LINE = re.compile(rf"(?:rs\.?|pkr)?\s*(\d[\d,]*(?:\.\d+)?)\s*(?:/|per|@|each)?\s*({_UNIT_WORDS})\b", re.I)


def rule_items(text: str) -> List[QuoteItem]:
    """No-AI reading: lines like 'Cement Lucky 50kg  1,450 / bag' or 'Sarya G60 12mm  262 per kg'."""
    out = []
    for line in text.splitlines():
        s = line.strip()
        if len(s) < 5 or not re.search(r"[a-zA-Z]{3}", s) or \
                re.search(r"quotation|invoice|phone|mobile|date|total|grand|ntn|strn", s, re.I):
            continue
        m = None
        for mm in _LINE.finditer(s):
            m = mm  # the last price-with-unit on the line is the rate
        if m:
            price = float(m.group(1).replace(",", ""))
            desc = s[: m.start()].strip(" -:|\t") or s
            if price > 0 and not re.fullmatch(r"[\d\s.,]+", desc):
                out.append(QuoteItem(desc[:120], price, m.group(2).lower()))
            continue
        nums = re.findall(r"\d[\d,]*(?:\.\d+)?", s)
        if nums and re.search(r"[a-zA-Z]{3}", s):  # table row without unit words: take the largest number as price
            vals = [float(x.replace(",", "")) for x in nums]
            desc = re.split(r"\s{2,}|\t|\|", s)[0]
            if max(vals) >= 5:
                out.append(QuoteItem(desc[:120], max(vals), ""))
    return out[:200]


def ai_items(text: str, images, keys) -> Tuple[List[QuoteItem], str]:
    from ai.groq_client import call_text_model, call_vision_model
    system = ("You read supplier quotations and bills for construction materials in Pakistan. Answer only JSON. "
              "Copy prices exactly; do not invent items.")
    prompt = ('Return {"items": [{"description": text, "unit": "bag|kg|ton|maund|brick|1000|cft|sft|sqm|rft|m|coil|'
              'gallon|litre|each|set|length", "price": number (price per that unit), "qty": number or null}]}.\n'
              + (f"TEXT:\n{text[:6000]}" if text else "Read the attached quotation image(s)."))
    try:
        raw = call_vision_model(getattr(keys, "groq", ""), system, prompt, images, max_tokens=1500, keys=keys) if images \
            else call_text_model(getattr(keys, "groq", ""), system, prompt, max_tokens=1500, json_mode=True, keys=keys)
        data = json.loads(re.sub(r"^```(?:json)?|```$", "", raw.strip(), flags=re.M))
    except Exception as exc:  # noqa: BLE001
        return [], f"AI could not read the quote ({str(exc)[:120]})"
    out = []
    for it in data.get("items") or []:
        try:
            out.append(QuoteItem(str(it.get("description", ""))[:120], float(it.get("price")), str(it.get("unit") or "").lower(),
                                 float(it["qty"]) if it.get("qty") not in (None, "") else None))
        except (TypeError, ValueError):
            continue
    return out, ""


# ---------------------------------------------------------------------------
# matching
# ---------------------------------------------------------------------------
def _tokens(s: str) -> set:
    s = s.lower()
    s = re.sub(r"(\d+)\s*mm", r"\1mm", s)
    s = re.sub(r"(\d+)\s*[x\*]\s*(\d+)\s*(?:in|\"|mm)?", r"\1x\2", s)
    for k, v in SYNONYMS.items():
        if re.search(rf"(?<![a-z]){re.escape(k)}(?![a-z])", s):
            s += " " + v
    return {t for t in re.findall(r"\d+x\d+|\d+/\d+|\d+mm|[a-z]{3,}|\d+", s)
            if t not in {"the", "and", "with", "for", "per", "down", "from"}}


def _is_spec(t: str) -> bool:
    return bool(re.search(r"\d", t))


def convert(price: float, quote_unit: str, db_unit: str) -> Optional[float]:
    qu = (quote_unit or "").lower().strip().rstrip(".")
    qb = UNIT_BASE.get(qu)
    db = DB_BASE.get(db_unit) or UNIT_BASE.get(db_unit.lower())
    if qb is None or db is None or qb[1] != db[1]:
        return None
    return price / qb[0] * db[0]


def match_items(items: List[QuoteItem], materials) -> List[QuoteMatch]:
    """materials: iterable of Material (knowledge base) - usually the project's purchase list."""
    mats = [(m, _tokens(f"{m.description} {m.specification} {m.category}")) for m in materials]
    out = []
    for it in items:
        qt = _tokens(it.text)
        best, best_s = None, 0.0
        for m, mt in mats:
            if not qt or not mt:
                continue
            common = qt & mt
            spec_q = {t for t in qt if _is_spec(t)}
            s = (len(common) + sum(1.0 for t in common if _is_spec(t))) / (len(qt) ** 0.5 * len(mt) ** 0.5)
            if spec_q and not (spec_q & mt):
                s *= 0.7  # the quote names a size (12mm, 24x24, 3/29) this material doesn't have
            if it.unit and DB_BASE.get(m.unit) and UNIT_BASE.get(it.unit) and DB_BASE[m.unit][1] == UNIT_BASE[it.unit][1]:
                s += 0.15  # compatible unit
            if s > best_s:
                best, best_s = m, s
        if best is None or best_s < 0.2:
            out.append(QuoteMatch(it, None, "", best_s, None, "", "no matching material"))
            continue
        conv = convert(it.price, it.unit, best.unit) if it.unit else (it.price if best_s > 0.45 else None)
        note = "" if conv is not None else f"unit '{it.unit or '?'}' can't be converted to {best.unit} - enter the rate"
        out.append(QuoteMatch(it, best.mat_id, best.description, round(best_s, 2),
                              round(conv, 2) if conv is not None else None, best.unit, note))
    return out
