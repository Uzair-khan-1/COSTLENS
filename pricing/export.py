"""Excel workbook for the cost: summary, priced materials, labour, cash flow and the government estimate.
Amounts are formulas (quantity x rate), so a rate typed in Excel updates the totals."""
from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pricing.costing import CONTRACTS, CostResult, pkr

NAVY, TEAL = "1B2F5B", "2E86DE"
FONT = "Arial"
THIN = Side(style="thin", color="D0D7E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0;[Red]-#,##0'
STATUS_FILL = {"live": "C6EFCE", "user": "DDEBF7", "indicative": "FFF2CC"}


def _hdr(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _c(ws, r, c, v, fmt=None, bold=False, fill=None, color=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name=FONT, size=10, bold=bold, color=color)
    cell.border = BORDER
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    return cell


def _title(ws, text, sub=""):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, size=9, color="666666")


def build_cost_workbook(cost: CostResult, project_name: str, gov=None) -> bytes:
    s = cost.settings
    wb = Workbook()
    ws = wb.active
    ws.title = "Cost_Summary"
    _title(ws, f"{project_name} - Cost estimate ({s.city})",
           f"Prepared {date.today():%d %b %Y}. Market prices for {s.city}; each rate shows its date and source. "
           "Quantities from the material take-off of the owner's scope.")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 46, 18, 50
    r = 4
    _c(ws, r, 1, "Contract type", bold=True)
    _c(ws, r, 2, CONTRACTS[s.contract][0])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 2
    _hdr(ws, r, ["Item", "Amount (Rs)", "Notes"], [46, 18, 50])
    ws.freeze_panes = None
    rows = [("Materials (all materials to buy, incl. wastage)", "=SUM(Materials_Priced!G:G)", "sheet Materials_Priced"),
            ("Labour (all BOQ work)", "=SUM(Labour_Priced!F:F)", "sheet Labour_Priced"),
            (f"Contractor profit & overhead @ {s.profit_pct:g}%", round(cost.profit), f"on {'labour' if s.contract == 'labour' else 'grey materials + labour' if s.contract == 'grey' else 'materials + labour'}")]
    rows += [(label, round(v), "optional owner cost") for label, v in cost.extras]
    rows += [(f"Contingency @ {s.contingency_pct:g}%", round(cost.contingency), ""),
             (f"Price rise during construction @ {s.escalation_pct_month:g}% per month", round(cost.escalation), "sheet Cash_Flow")]
    first = r + 1
    for label, v, note in rows:
        r += 1
        _c(ws, r, 1, label)
        _c(ws, r, 2, v, MONEY)
        _c(ws, r, 3, note, color="777777")
    r += 1
    _c(ws, r, 1, "TOTAL COST", bold=True, fill="E2EFDA")
    _c(ws, r, 2, f"=SUM(B{first}:B{r - 1})", MONEY, bold=True, fill="E2EFDA")
    _c(ws, r, 3, pkr(cost.total), bold=True, fill="E2EFDA")
    r += 1
    _c(ws, r, 1, f"Cost per sq ft of covered area ({cost.covered_sft:,.0f} sq ft)")
    _c(ws, r, 2, f"=B{r - 1}/{cost.covered_sft:.2f}" if cost.covered_sft else 0, MONEY)
    r += 1
    _c(ws, r, 1, "Likely range (low - high rates)")
    _c(ws, r, 2, f"{pkr(cost.low)} - {pkr(cost.high)}")
    r += 1
    _c(ws, r, 1, "Paid to the contractor / bought by you")
    _c(ws, r, 2, f"{pkr(cost.contractor_part)} / {pkr(cost.owner_buys)}")
    if gov is not None:
        r += 2
        _c(ws, r, 1, "Government estimate (MRS / CSR rates, for comparison)", bold=True)
        _c(ws, r, 2, round(gov.total), MONEY, bold=True)
        _c(ws, r, 3, f"sheet Government_Estimate; {gov.mapped_share:.0%} priced from MRS/CSR items")
    r += 2
    _c(ws, r, 1, "Cost by trade", bold=True, color=NAVY)
    r += 1
    _hdr(ws, r, ["Trade", "Materials (Rs)", "Labour (Rs)"], [46, 18, 18])
    ws.freeze_panes = None
    for t, m, lab in cost.by_trade():
        r += 1
        _c(ws, r, 1, t)
        _c(ws, r, 2, round(m), MONEY)
        _c(ws, r, 3, round(lab), MONEY)
    r += 2
    for n in cost.notes:
        _c(ws, r, 1, n, color="9C5700")
        r += 1

    # ---- materials
    wm = wb.create_sheet("Materials_Priced")
    _title(wm, "Materials - quantities x rates", "Change a rate (column F) to your supplier's price - the amount updates. "
           "Green = live market price, blue = your quote, yellow = indicative rate book.")
    _hdr(wm, 4, ["Code", "Material", "Trade", "Unit", "Qty (incl. wastage)", "Rate (Rs)", "Amount (Rs)", "Buy", "Rate status",
                 "Rate date", "Source"], [10, 44, 22, 7, 14, 12, 14, 14, 11, 11, 40])
    r = 4
    for ln in [x for x in cost.lines if x.kind == "material"]:
        r += 1
        _c(wm, r, 1, ln.code)
        _c(wm, r, 2, ln.name)
        _c(wm, r, 3, ln.trade)
        _c(wm, r, 4, ln.unit)
        _c(wm, r, 5, round(ln.qty, 2), "#,##0.00")
        _c(wm, r, 6, round(ln.rate, 2), "#,##0.00", fill=STATUS_FILL.get(ln.status)).font = Font(name=FONT, size=10, color="0000FF")
        _c(wm, r, 7, f"=E{r}*F{r}", MONEY)
        _c(wm, r, 8, ln.buy_qty)
        _c(wm, r, 9, ln.status, fill=STATUS_FILL.get(ln.status))
        _c(wm, r, 10, ln.as_of)
        _c(wm, r, 11, ln.source)
    wm.auto_filter.ref = f"A4:K{r}"

    # ---- labour
    wl = wb.create_sheet("Labour_Priced")
    _title(wl, "Labour - BOQ work x labour (thekedar) rates", "Labour-only rates per unit of work for the city.")
    _hdr(wl, 4, ["WI", "Work", "Unit", "Quantity", "Rate (Rs)", "Amount (Rs)", "Rate status", "Source"],
         [11, 50, 7, 12, 11, 14, 11, 40])
    r = 4
    for ln in [x for x in cost.lines if x.kind == "labour"]:
        r += 1
        _c(wl, r, 1, ln.code)
        _c(wl, r, 2, ln.name)
        _c(wl, r, 3, ln.unit)
        _c(wl, r, 4, round(ln.qty, 2), "#,##0.00")
        _c(wl, r, 5, round(ln.rate, 2), "#,##0.00", fill=STATUS_FILL.get(ln.status)).font = Font(name=FONT, size=10, color="0000FF")
        _c(wl, r, 6, f"=D{r}*E{r}", MONEY)
        _c(wl, r, 7, ln.status, fill=STATUS_FILL.get(ln.status))
        _c(wl, r, 8, ln.source)
    wl.auto_filter.ref = f"A4:H{r}"

    # ---- cash flow
    if cost.cashflow:
        wc = wb.create_sheet("Cash_Flow")
        _title(wc, "Money needed each month", "From the construction schedule: materials bought when first needed, "
               "labour paid as work is done, price rise per month added.")
        cols = ["Month", "Month of", "Materials", "Labour", "Contractor, extras & contingency", "Price rise (escalation)",
                "Total", "Cumulative"]
        _hdr(wc, 4, cols, [8, 12, 14, 14, 18, 16, 14, 16])
        r = 4
        for m in cost.cashflow:
            r += 1
            for j, k in enumerate(cols, 1):
                v = m.get(k)
                _c(wc, r, j, round(v) if isinstance(v, float) else v,
                   MONEY if isinstance(v, float) else ("MMM YYYY" if k == "Month of" else None))

    # ---- government estimate
    if gov is not None:
        wg = wb.create_sheet("Government_Estimate")
        _title(wg, f"Government estimate - {s.city}", f"Rates from {gov.book} {gov.edition} (Balochistan CSR-2026 for Quetta). "
               "Items without a government item are priced at market cost (see Source).")
        _hdr(wg, 4, ["WI", "Description", "Unit", "Quantity", "Rate (Rs)", "Amount (Rs)", "Source", "Item"],
             [11, 50, 7, 12, 12, 14, 40, 22])
        r = 4
        for ln in gov.lines:
            r += 1
            _c(wg, r, 1, ln["WI_ID"])
            _c(wg, r, 2, ln["Description"])
            _c(wg, r, 3, ln["Unit"])
            _c(wg, r, 4, ln["Quantity"], "#,##0.00")
            _c(wg, r, 5, ln["Rate"], "#,##0.00")
            _c(wg, r, 6, f"=D{r}*E{r}", MONEY)
            _c(wg, r, 7, ln["Source"])
            _c(wg, r, 8, ln["Item"])
        sub = r
        r += 1
        _c(wg, r, 2, "Sub-total", bold=True)
        _c(wg, r, 6, f"=SUM(F5:F{sub})", MONEY, bold=True)
        base = r
        for label, v in gov.extras:
            r += 1
            _c(wg, r, 2, label)
            pct = float(label.split("@")[1].strip(" %")) if "@" in label else 0
            _c(wg, r, 6, f"=F{base}*{pct / 100:g}", MONEY)
        r += 1
        _c(wg, r, 2, "GRAND TOTAL", bold=True, fill="E2EFDA")
        _c(wg, r, 6, f"=SUM(F{base}:F{r - 1})", MONEY, bold=True, fill="E2EFDA")
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def cost_pdf(cost: CostResult, project_name: str, gov=None) -> Optional[bytes]:
    """One-page cost summary PDF (ReportLab)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    s = cost.settings
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm)
    ss = getSampleStyleSheet()
    small = ss["BodyText"].clone("s", fontSize=8.5, leading=11)
    logo = pdf_logo()
    story = ([logo] if logo else []) + [Paragraph(f"<b>{project_name}</b> - Cost estimate", ss["Title"]),
             Paragraph(f"{s.city} \u00b7 {CONTRACTS[s.contract][0]} \u00b7 {date.today():%d %b %Y} \u00b7 covered area "
                       f"{cost.covered_sft:,.0f} sq ft", small), Spacer(1, 4 * mm)]
    rows = [["", "Amount"], ["Materials", pkr(cost.materials)], ["Labour", pkr(cost.labour)],
            [f"Contractor profit & overhead ({s.profit_pct:g}%)", pkr(cost.profit)]]
    rows += [[l, pkr(v)] for l, v in cost.extras]
    rows += [[f"Contingency ({s.contingency_pct:g}%)", pkr(cost.contingency)],
             [f"Price rise during construction ({s.escalation_pct_month:g}%/month)", pkr(cost.escalation)],
             ["TOTAL", pkr(cost.total)], ["Per sq ft (covered)", f"Rs {cost.per_sft:,.0f}"],
             ["Likely range", f"{pkr(cost.low)} - {pkr(cost.high)}"]]
    if gov is not None:
        rows.append(["Government estimate (MRS/CSR, for comparison)", pkr(gov.total)])
    t = Table(rows, colWidths=[115 * mm, 55 * mm])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B2F5B")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("FONTSIZE", (0, 0), (-1, -1), 9.5), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D7E2")),
                           ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                           ("FONTNAME", (0, len(rows) - (4 if gov is not None else 3)), (-1, len(rows) - (4 if gov is not None else 3)), "Helvetica-Bold"),
                           ("BACKGROUND", (0, len(rows) - (4 if gov is not None else 3)), (-1, len(rows) - (4 if gov is not None else 3)),
                            colors.HexColor("#E2EFDA"))]))
    story += [t, Spacer(1, 5 * mm), Paragraph("<b>Largest costs by trade</b>", ss["Heading4"])]
    tr = [["Trade", "Materials", "Labour"]] + [[n, pkr(a), pkr(b)] for n, a, b in cost.by_trade()[:12]]
    t2 = Table(tr, colWidths=[90 * mm, 40 * mm, 40 * mm])
    t2.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E86DE")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("FONTSIZE", (0, 0), (-1, -1), 9), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D7E2")),
                            ("ALIGN", (1, 0), (-1, -1), "RIGHT")]))
    story += [t2, Spacer(1, 4 * mm)]
    if cost.cashflow:
        cf = [["Month", "Money needed", "Cumulative"]] + [[f"{m['Month of']:%b %Y}", pkr(m["Total"]), pkr(m["Cumulative"])]
                                                           for m in cost.cashflow if m["Total"] > 0]
        t3 = Table(cf, colWidths=[40 * mm, 45 * mm, 45 * mm])
        t3.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B2F5B")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                ("FONTSIZE", (0, 0), (-1, -1), 9), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D7E2")),
                                ("ALIGN", (1, 0), (-1, -1), "RIGHT")]))
        story += [Paragraph("<b>Money needed each month</b>", ss["Heading4"]), t3]
    story += [Spacer(1, 4 * mm), Paragraph("Market prices are indications (dealers vary 5-10%); get 2-3 quotes before buying. "
                                           "Not a certified QS estimate.", small)]
    doc.build(story)
    return buf.getvalue()


def pdf_logo(width_mm: float = 52.0):
    """The CostLens logo as a ReportLab flowable (None if the file is missing)."""
    try:
        import config
        from reportlab.lib.units import mm
        from reportlab.platypus import Image
        if config.LOGO_HORIZONTAL_PATH.exists():
            img = Image(str(config.LOGO_HORIZONTAL_PATH), width=width_mm * mm, height=width_mm * mm * 512 / 1700)
            img.hAlign = "LEFT"
            return img
    except Exception:  # noqa: BLE001
        pass
    return None
